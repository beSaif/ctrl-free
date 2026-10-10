"""Export tests on a synthetic tender built from the 22 verified examples."""
import io
import shutil
import subprocess

import openpyxl
import pymupdf
import pytest

from app import export as E
from app import pricing as P
from tests.test_pricing import example_tender


def tender_with_extras() -> dict:
    """The 22 examples plus a heading, a line without quantity and an out-of-scope line."""
    t = example_tender()
    sys0 = t["positions"][0]["system"]
    t["positions"] += [
        {"lv_id": "LV0023", "system": sys0, "page": 81, "kind": "subtotal", "title_de": "Total Armaturen C"},
        {"lv_id": "LV0024", "system": sys0, "page": 81, "kind": "item", "title_de": "Regulierklappe DN160",
         "quantity": None, "unit": "pc", "match": {"product_id": "P-DAM-001", "install_code": "DAMPER", "confidence": 0.9}},
        {"lv_id": "LV0025", "system": sys0, "page": 82, "kind": "item", "title_de": "Kernbohrung bauseits",
         "quantity": 3, "unit": "pc", "match": {"product_id": "P-SUP-001", "install_code": "SUPPORT_SET", "confidence": 0.9},
         "status": "out_of_scope", "status_reason": "bauseits (by the client)"},
    ]
    return t


def modified_tender() -> dict:
    """Overrides, a supplier quote and every switch on, to exercise all formula paths."""
    t = tender_with_extras()
    by_id = {p["lv_id"]: p for p in t["positions"]}
    by_id["LV0001"]["overrides"] = {"hours": 6.5}
    by_id["LV0002"]["overrides"] = {"unit_price": 90.0, "quantity": 7}
    by_id["LV0006"]["supplier_quote"] = {"unit_price": 2450.0, "source": "Supplier reply.pdf p.1"}
    s = t["pricing"]
    s.update(chosen="Target", include_small_material=True, include_supplier_contingency=True, overhead_pct=9.0)
    s["project_costs"] = {k: True for k in s["project_costs"]}
    s["factors"]["site_efficiency"] = 0.85
    s["competition"] = -0.10
    s["notes"]["competition"] = "heard 4 firms took the documents"
    s["changed_by"] = {"competition": {"by": "You", "at": "2026-09-10T09:12:00"},
                       "overhead_pct": {"by": "You", "at": "2026-09-10T09:13:00"}}
    return t


def _totals(t):
    return P.project_totals(t["positions"], t["pricing"], t["systems"])


# ---------------------------------------------------------------- xlsx

def test_xlsx_structure_and_formulas():
    t = tender_with_extras()
    tot = _totals(t)
    wb = openpyxl.load_workbook(io.BytesIO(E.build_xlsx(t, tot, t["pricing"])))
    assert wb.sheetnames == ["Offer", "Summary", "Catalog", "Install rules", "Parameters"]
    ws = wb["Offer"]
    hdr = {c.value: c.column_letter for c in ws[1]}
    for col in ("LV id", "System", "Section", "Page", "Tender text", "Qty", "Unit", "Product id", "Product description",
                "Net unit price", "Waste %", "Material", "Base h/unit", "Factor product", "Hours", "Rate", "Labour",
                "Direct", "Status", "Source page"):
        assert col in hdr, col
    rows = {ws[f"{hdr['LV id']}{r}"].value: r for r in range(2, ws.max_row + 1) if ws[f"{hdr['LV id']}{r}"].value}
    assert len(rows) == 25
    r = rows["LV0001"]
    assert ws[f"{hdr['Material']}{r}"].value.startswith("=ROUND(")
    assert ws[f"{hdr['Waste %']}{r}"].value == "=waste_pct"
    assert ws[f"{hdr['Factor product']}{r}"].value == "=1/f_site_efficiency*f_mounting_height_school*f_school_logistics"
    assert ws[f"{hdr['Hours']}{r}"].value.startswith("=")
    assert ws[f"{hdr['Labour']}{r}"].value.startswith("=ROUND(")
    assert ws[f"{hdr['Direct']}{r}"].value == f"={hdr['Material']}{r}+{hdr['Labour']}{r}"
    assert "Catalog!" in ws[f"{hdr['Net unit price']}{r}"].value
    assert "'Install rules'!" in ws[f"{hdr['Rate']}{r}"].value
    assert ws[f"{hdr['Page']}{r}"].comment.text == "Tender p. 78"
    assert ws[f"{hdr['Product id']}{r}"].comment.text == "catalog.csv P-AIR-001"
    assert "productivity.csv DIFFUSER" in ws[f"{hdr['Install rule']}{r}"].comment.text
    # lines without a price: status shown, amounts blank
    for lv, st in (("LV0023", "non_priceable"), ("LV0024", "insufficient_information"), ("LV0025", "out_of_scope")):
        r = rows[lv]
        assert st in ws[f"{hdr['Status']}{r}"].value
        assert ws[f"{hdr['Direct']}{r}"].value is None and ws[f"{hdr['Material']}{r}"].value is None
    # Summary formulas and named parameters
    sm = wb["Summary"]
    formulas = [c.value for row in sm.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=SUM(Offer!") for f in formulas)
    assert sum("CEILING(" in f for f in formulas) == 5  # standard price + 4 options
    assert any("transport_pct" in f for f in formulas) and any("overhead_pct" in f for f in formulas)
    for n in ("waste_pct", "transport_pct", "overhead_pct", "f_site_efficiency", "margin_Standard", "margin_Target",
              "pc_engineering_on", "total_direct", "total_cost"):
        assert n in wb.defined_names, n
    assert wb["Catalog"].max_row > 40 and wb["Install rules"].max_row > 20


def _recalc(tmp_path, files: dict[str, bytes]) -> dict[str, openpyxl.Workbook]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pytest.skip("LibreOffice not installed; formula recalculation not checked")
    paths = []
    for name, data in files.items():
        p = tmp_path / f"{name}.xlsx"
        p.write_bytes(data)
        paths.append(str(p))
    out = tmp_path / "recalc"
    try:
        subprocess.run([soffice, f"-env:UserInstallation=file://{tmp_path}/lo_profile", "--headless", "--calc",
                        "--convert-to", "xlsx", "--outdir", str(out), *paths],
                       check=True, capture_output=True, timeout=240)
    except (subprocess.SubprocessError, OSError) as e:
        pytest.skip(f"LibreOffice could not recalculate: {e}")
    return {n: openpyxl.load_workbook(out / f"{n}.xlsx", data_only=True) for n in files}


def _summary_values(wb) -> dict:
    ws = wb["Summary"]
    vals = {}
    for row in ws.iter_rows(values_only=True):
        if row and isinstance(row[0], str):
            vals[row[0].split("  ·")[0]] = row[1:]
    return vals


def test_xlsx_recalculates_to_python_totals(tmp_path):
    t1 = tender_with_extras()
    t2 = modified_tender()
    tot1, tot2 = _totals(t1), _totals(t2)
    wbs = _recalc(tmp_path, {"default": E.build_xlsx(t1, tot1, t1["pricing"]),
                             "modified": E.build_xlsx(t2, tot2, t2["pricing"])})
    for name, tot in (("default", tot1), ("modified", tot2)):
        v = _summary_values(wbs[name])
        direct = next(val for k, val in v.items() if k.startswith("Direct cost"))[0]
        assert direct == pytest.approx(tot["direct"], abs=0.005), name
        assert v["Material"][0] == pytest.approx(tot["material"], abs=0.005)
        assert v["Labour"][0] == pytest.approx(tot["labour"], abs=0.005)
        assert v["+ Project costs switched on"][0] == pytest.approx(tot["project_costs"], abs=0.005)
        assert v["Cost"][0] == pytest.approx(tot["cost"], abs=0.02), name
        assert v["Standard price (standard margin, rounded up)"][0] == tot["standard_price"]
        for sc in tot["scenarios"]:
            row = v[sc["name"]]
            assert row[1] == sc["bid"], (name, sc["name"])
            assert round(row[2], 2) == sc["win"]
            assert row[4] == pytest.approx(sc["expected"], abs=0.05)
    assert tot1["direct"] == pytest.approx(372700.12 - 0, abs=0.005)  # extras add nothing priced
    assert tot2["project_costs"] == 115700 and tot2["small_material"] > 0 and tot2["supplier_contingency"] > 0
    assert _summary_values(wbs["modified"])["Offer price"][0] == tot2["offer_price"]


# ---------------------------------------------------------------- pdf

def test_report_pdf():
    t = modified_tender()
    tot = _totals(t)
    pdf = E.build_report_pdf(t, tot, t["pricing"], E.open_points(t, tot))
    assert pdf[:5] == b"%PDF-"
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    text = " ".join(" ".join(page.get_text() for page in doc).split())
    assert doc.page_count >= 2
    assert "Pricing report" in text
    assert f"Offer price CHF {P.swiss(tot['offer_price'], 0)}" in text
    assert "How the cost adds up" in text and "Project costs switched on" in text
    assert "Your call · You · 10 Sep 2026" in text and "heard 4 firms took the documents" in text
    assert "Project parameters" in text
    assert "Open points before you submit" in text and "LV0018" in text and "Assumptions" in text
    assert "Ask supplier (supplier_quote_required)" in text
    assert "÷ 0.85" in text


def test_report_pdf_no_offer_chosen():
    t = example_tender()
    tot = _totals(t)
    text = pymupdf.open(stream=E.build_report_pdf(t, tot, t["pricing"], E.open_points(t, tot)), filetype="pdf")[0].get_text()
    assert "Offer price not chosen yet" in text and "439'000" in text


# ---------------------------------------------------------------- SIA

def test_sia_file():
    t = tender_with_extras()
    t["pricing"]["chosen"] = "Target"
    tot = _totals(t)
    text = E.build_sia_file(t, tot).decode("utf-8")
    lines = text.splitlines()
    assert lines[0].startswith("# ") and "simplified" in text and "SIA 451" in text and "next step" in text
    header = next(l for l in lines if l.startswith("Pos;"))
    assert header.split(";")[6:10] == ["Menge", "Einheit", "Einheitspreis CHF", "Betrag CHF"]
    rows = [l.split(";") for l in lines if l[:1].isdigit()]
    assert len(rows) == 24  # subtotal line skipped
    lv = {r[1]: r for r in rows}
    assert lv["LV0001"][6:8] == ["5", "St"] and float(lv["LV0001"][9]) == pytest.approx(5 * float(lv["LV0001"][8]))
    assert lv["LV0024"][9] == "" and lv["LV0025"][9] == "" and lv["LV0025"][10] == "out_of_scope"
    total = sum(float(r[9]) for r in rows if r[9])
    tot_line = next(l for l in lines if l.startswith("TOTAL;")).split(";")
    assert float(tot_line[9]) == pytest.approx(total, abs=0.01)
    assert abs(total - 432000) < 25  # unit-price rounding only
    assert next(l for l in lines if l.startswith("OFFER;")).split(";")[9] == "432000.00"


# ---------------------------------------------------------------- open points

def test_open_points():
    t = example_tender()
    tot = _totals(t)
    pts = {p["key"]: p for p in E.open_points(t, tot)}
    assert pts["lines"]["title"] == "16 of 22 lines priced from the catalogue"
    assert pts["budget"]["title"] == "6 lines on budget prices"
    assert pts["budget"]["detail"] == "CHF 271'261 of direct cost (73%). No supplier quotes entered yet."
    assert pts["instructions"]["title"] == "2 of 3 instructions not confirmed"
    assert pts["instructions"]["detail"] == "1 adds cost, 1 is a risk"
    assert pts["project_costs"]["title"] == "Project costs not included" and "115'700" in pts["project_costs"]["detail"]
    assert pts["offer"]["title"] == "Offer price not chosen" and not pts["offer"]["ok"]
    t["pricing"]["chosen"] = "Target"
    tot = _totals(t)
    pts = {p["key"]: p for p in E.open_points(t, tot)}
    assert pts["offer"]["ok"] and pts["offer"]["title"] == "Offer price CHF 432'000"
    assert pts["offer"]["detail"] == "Target, 4.0% margin · 63% win chance"
    t2 = modified_tender()
    pts2 = {p["key"]: p for p in E.open_points(t2, _totals(t2))}
    assert pts2["project_costs"]["ok"] and "1 supplier quote entered" in pts2["budget"]["detail"]
