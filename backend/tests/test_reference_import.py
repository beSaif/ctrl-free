"""Catalogue, install-rule and price-list import from any table format (app.tables):
xlsx, ods, JSON, semicolon CSV with German headers and Swiss numbers, UTF-16, and PDF tables.

Every test works on a copy of the reference files (REFERENCE_DIR is patched), never on data/reference."""
from __future__ import annotations

import csv
import io
import json

import pytest

from app import config, price_lists as PL, reference as R, tables

GERMAN = {"product_id": "Art.-Nr.", "category": "Kategorie", "subcategory": "Unterkategorie",
          "manufacturer": "Hersteller", "model": "Typ", "description": "Bezeichnung", "unit": "ME",
          "list_price_chf": "Listenpreis CHF", "discount_pct": "Rabatt %", "net_cost_chf": "Nettopreis CHF",
          "install_code": "Montagecode", "source_type": "Quelle", "supplier_quote_required": "Offerte erforderlich"}


@pytest.fixture
def ref_dir(tmp_path, monkeypatch):
    for f in R.FILES.values():
        (tmp_path / f).write_bytes((config.REFERENCE_DIR / f).read_bytes())
    monkeypatch.setattr(R, "REFERENCE_DIR", tmp_path)
    R.clear_cache()
    yield tmp_path
    monkeypatch.undo()
    R.clear_cache()


def _orig_rows() -> list[dict]:
    return list(csv.DictReader(io.StringIO((config.REFERENCE_DIR / "catalog.csv").read_text("utf-8"))))


def _expected() -> dict:
    return R._parse_catalog((config.REFERENCE_DIR / "catalog.csv").read_text("utf-8"))


def _german_rows(swiss: bool = False) -> list[list]:
    """The demo catalogue as a German export: Stk, ja/nein, prices like 1'487.50 / 1.487,50 / CHF 74.80."""
    out = [[GERMAN[c] for c in R.CATALOG_COLUMNS]]
    styles = [lambda x: f"{x:,.2f}".replace(",", "'"), lambda x: f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."),
              lambda x: f"CHF {x:.2f}"]
    for i, r in enumerate(_orig_rows()):
        row = []
        for c in R.CATALOG_COLUMNS:
            v = r[c]
            if c == "unit":
                v = {"pc": "Stk", "m": "lfm"}.get(v, v)
            elif c == "supplier_quote_required":
                v = "ja" if v == "true" else "nein"
            elif c == "net_cost_chf" and swiss:
                v = styles[i % 3](float(v))
            row.append(v)
        out.append(row)
    return out


def _same_catalog():
    got, exp = R.load_catalog(), _expected()
    assert set(got) == set(exp)
    for pid, e in exp.items():
        assert got[pid] == e, pid


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(rows: list[list], title: str = "Katalog Equans 2026") -> bytes:
    """A table with ruling lines over several pages, the header repeated on each page."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = io.BytesIO()
    t = Table(rows, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTSIZE", (0, 0), (-1, -1), 6)]))
    SimpleDocTemplate(buf, pagesize=landscape(A4)).build(
        [Paragraph(title, getSampleStyleSheet()["Title"]), Spacer(1, 8), t])
    return buf.getvalue()


# ---------------------------------------------------------------- reading

def test_formats_are_detected_by_content():
    assert tables.detect_format(b"%PDF-1.7 ...", "x.csv") == "pdf"
    assert tables.detect_format(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "export") == "xls"
    assert tables.detect_format(_xlsx({"S": [["a"]]}), "upload.bin") == "xlsx"
    assert tables.detect_format(b'[{"a": 1}]', "data.txt") == "json"
    assert tables.detect_format(b"a;b\n1;2\n", "x.tsv") == "text"
    with pytest.raises(ValueError, match="old Excel"):
        tables.read_table(b"\xd0\xcf\x11\xe0 not really a workbook", "x.xls")


def test_numbers_and_text_encodings():
    assert tables.num("1'234.50") == 1234.5 and tables.num("1.234,50") == 1234.5
    assert tables.num("CHF 12.00") == 12 and tables.num("12.-") == 12 and tables.num("1.234.567") == 1234567
    rows = tables.read_table("Bezeichnung\tPreis\nZuluftventil\t74,80\n".encode("utf-16"), "liste.txt")
    assert rows == [["Bezeichnung", "Preis"], ["Zuluftventil", "74,80"]]
    rows = tables.read_table("Größe|Preis\nKlappe|1'487.50\n".encode("cp1252"), "liste.csv")
    assert rows[0] == ["Größe", "Preis"] and rows[1][1] == "1'487.50"


# ---------------------------------------------------------------- catalogue

def test_catalogue_from_semicolon_csv_with_german_headers_and_swiss_numbers(ref_dir):
    rows = _german_rows(swiss=True)
    text = "Katalog Equans, Stand März 2026\n\n" + "\n".join(";".join(r) for r in rows) + "\n"
    entry = R.import_reference("catalog", text.encode("cp1252"), filename="Katalog März.csv")
    assert entry["format"] == "csv" and entry["rows"] == len(rows) - 1
    _same_catalog()
    assert (ref_dir / "catalog.csv").read_text("utf-8").splitlines()[0] == ",".join(R.CATALOG_COLUMNS)


def test_catalogue_from_xlsx_picks_the_sheet_with_the_header(ref_dir):
    data = _xlsx({"Deckblatt": [["Equans Produktkatalog"], ["Stand 2026"]],
                  "Artikel": [["Katalog 2026"], []] + _german_rows(swiss=True)})
    entry = R.import_reference("catalog", data, filename="katalog.xlsx")
    assert entry["format"] == "xlsx" and entry["original_name"] == "katalog.xlsx"
    _same_catalog()
    upload = ref_dir / "history" / entry["upload"]
    assert upload.read_bytes() == data and upload.name.endswith("katalog.xlsx")
    v = R.reference_versions()[-1]
    assert v["original_name"] == "katalog.xlsx" and v["format"] == "xlsx" and v["previous"]


def test_catalogue_from_ods(ref_dir):
    import pandas as pd
    rows = _german_rows()
    buf = io.BytesIO()
    pd.DataFrame(rows[1:], columns=rows[0]).to_excel(buf, engine="odf", index=False)
    entry = R.import_reference("catalog", buf.getvalue(), filename="katalog.ods")
    assert entry["format"] == "ods"
    _same_catalog()


def test_catalogue_from_json_without_source_columns(ref_dir):
    items = [{k: v for k, v in r.items() if k not in ("source_type", "supplier_quote_required")} for r in _orig_rows()]
    for it in items:
        it["net_cost_chf"] = float(it["net_cost_chf"])
    entry = R.import_reference("catalog", json.dumps({"catalog": items}).encode(), filename="catalog.json")
    assert entry["format"] == "json" and entry["defaulted_columns"] == ["source_type", "supplier_quote_required"]
    got = R.load_catalog()
    assert got["P-AIR-001"]["net_cost_chf"] == 74.8 and got["P-AIR-001"]["source_type"] == "catalog"
    assert got["P-SIL-001"]["supplier_quote_required"] is False


def test_catalogue_from_pdf_table_over_several_pages(ref_dir):
    data = _pdf(_german_rows(swiss=True))
    import pymupdf
    assert pymupdf.open(stream=data).page_count >= 2
    entry = R.import_reference("catalog", data, filename="Katalog.pdf")
    assert entry["format"] == "pdf" and entry["rows"] == len(_orig_rows())
    _same_catalog()                       # FIRE_DAMPER etc. keep their underscores


def test_pdf_without_table_and_missing_columns_are_plain_errors(ref_dir):
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 720, "Sehr geehrte Damen und Herren, anbei unser Katalog.")
    c.save()
    with pytest.raises(ValueError, match="No table found in this PDF. Export the catalogue as Excel or CSV"):
        R.import_reference("catalog", buf.getvalue(), filename="brief.pdf")
    with pytest.raises(ValueError, match=r"missing columns: category, install_code\. Headers found: 'Art\.-Nr\.', 'Bezeichnung'"):
        R.import_reference("catalog", "Art.-Nr.;Bezeichnung;ME;Preis\nX1;Ventil;Stk;12,50\n".encode(), filename="x.csv")
    with pytest.raises(ValueError, match="Row 3: product_id is empty"):
        R.import_reference("catalog", _xlsx({"A": [_german_rows()[0], _german_rows()[1], ["", "Air terminal", "", "", "", "x", "Stk", "", "", "12", "DIFFUSER", "", ""]]}),
                           filename="x.xlsx")
    assert R.reference_versions() == []   # nothing saved


def test_install_rules_from_xlsx_with_german_headers(ref_dir):
    orig = list(csv.DictReader(io.StringIO((config.REFERENCE_DIR / "productivity.csv").read_text("utf-8"))))
    head = ["Montagecode", "Arbeitsart", "Einheit", "Std/ME", "Rolle", "Stundenansatz CHF", "Faktoren"]
    rows = [["Montageregeln Equans"], head] + [
        [r["install_code"], r["work_type"], r["unit"], r["base_hours_per_unit"].replace(".", ","), r["labour_role"],
         float(r["default_hourly_rate_chf"]), r["factor_keys"].replace("|", ", ")] for r in orig]
    entry = R.import_reference("rules", _xlsx({"Regeln": rows}), filename="regeln.xlsx")
    assert entry["kind"] == "rules" and entry["rows"] == len(orig)
    exp = R._parse_rules((config.REFERENCE_DIR / "productivity.csv").read_text("utf-8"))
    assert R.load_rules() == exp


def test_api_import_passes_the_file_name_and_turns_errors_into_400(ref_dir):
    from fastapi.testclient import TestClient
    from app.main import app
    client = TestClient(app)
    data = _xlsx({"Artikel": _german_rows()})
    r = client.post("/api/reference/catalog", files={"file": ("Katalog 2026.xlsx", data)})
    assert r.status_code == 200, r.text
    assert r.json()["version"]["original_name"] == "Katalog 2026.xlsx"
    r = client.post("/api/reference/catalog", files={"file": ("x.csv", b"foo;bar\n1;2\n")})
    assert r.status_code == 400 and "missing columns" in r.json()["detail"]


# ---------------------------------------------------------------- price lists

def test_price_list_from_pdf_and_ods(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PRICE_LISTS_DIR", tmp_path / "price_lists")
    rows = [["Artikel-Nr.", "Hersteller", "Typ", "Bezeichnung", "Einheit", "Bruttopreis CHF", "Rabatt %", "Netto CHF"],
            ["DV1", "DemoVent", "DV-Z-100", "Zuluftventil DN100", "Stk", "85.00", "12", "74.80"],
            ["DV2", "DemoVent", "DV-Z-125", "Zuluftventil DN125", "Stk", "98.00", "12", "86.24"],
            ["DV3", "DemoVent", "DV-Z-999", "Unbekannt", "Stk", "1'000.00", "10", "900.00"]]
    p = PL.preview(_pdf(rows, "Preisliste DemoVent AG 2026"), "demovent.pdf")
    assert p["mapping"]["model"] == "Typ" and p["mapping"]["net_price"] == "Netto CHF"
    assert p["counts"]["matched"] == 2 and p["counts"]["unmatched"] == 1
    entry = PL.create(_pdf(rows), "demovent.pdf", None, {"name": "DemoVent", "kind": "supplier",
                                                          "valid_from": "2026-01-01", "supplier": "DemoVent"},
                      activate=False)
    assert entry["format"] == "pdf" and entry["file"].endswith(".pdf")

    import pandas as pd
    buf = io.BytesIO()
    pd.DataFrame(rows[1:], columns=rows[0]).to_excel(buf, engine="odf", index=False)
    assert PL.preview(buf.getvalue(), "demovent.ods")["counts"]["matched"] == 2
