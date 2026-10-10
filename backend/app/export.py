"""Exports: offer workbook (live formulas), pricing report PDF, simplified SIA-like file, open points.

All numbers come from ``app.pricing`` (positions are re-priced here, never trusted
from storage). The workbook re-derives every amount with formulas that point at
its own Catalog, Install rules and Parameters sheets, so an estimator can change
a parameter in Excel and see the offer move.
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from app import pricing as P
from app import reference as ref

INK, MUTED, RULE, PANEL, ACCENT, ATTN = "121212", "5A5A57", "DEDEDB", "F6F6F4", "1F4FD8", "C2410C"

COMPETITION_WORDS = {0.05: "Light", 0.0: "Normal", -0.05: "Strong", -0.10: "Very strong"}
APPETITE_WORDS = {-0.05: "Not a priority", 0.0: "Neutral", 0.05: "We want it", 0.10: "Must win"}
SENSITIVITY_WORDS = {0.03: "Low", 0.05: "Normal", 0.08: "High"}
STATUS_HELP = {
    "calculated": "Catalogue product matched with confidence, quantity known.",
    "assumption_applied": "Matched, but an attribute was assumed or the match is only close.",
    "supplier_quote_required": "Supplier package or named offer: budget price until a quote is entered.",
    "review_required": "No confident match; the estimator must pick the product.",
    "insufficient_information": "No quantity (or quantity 0) in the tender.",
    "out_of_scope": "Other trade, excluded lot or 'bauseits'; not in the price.",
    "non_priceable": "Heading, total or pure text; nothing to price.",
}
SIA_EXTENSION = "csv"  # simplified semicolon CSV; real SIA 451 is the production next step
SIA_MEDIA_TYPE = "text/csv; charset=utf-8"
SIA_UNITS = {"pc": "St", "m": "m", "m2": "m2", "m3": "m3", "kg": "kg", "h": "h", "LS": "gl", "system": "Anl"}


# ---------------------------------------------------------------- shared

def _settings(tender: dict, settings: dict | None = None) -> dict:
    return ref.merge_settings(settings if settings is not None else tender.get("pricing"))


def _evaluate(tender: dict, settings: dict) -> list[dict]:
    """Every position with fresh pricing/status/needs_review."""
    return P.evaluate_positions(tender.get("positions") or [], settings, tender.get("systems"),
                                tender.get("profile") or "school")


def _page(p: dict):
    return p.get("page") or (p.get("source") or {}).get("page")


def _title(p: dict) -> str:
    return p.get("title_de") or (p.get("text_de") or "").split("\n")[0] or p.get("title_en") or ""


def _offer(totals: dict) -> tuple[float, str]:
    """Offer price and how it was picked (chosen option, else the standard price)."""
    if totals.get("offer_price"):
        return float(totals["offer_price"]), f"{totals['offer']['name']} option"
    return float(totals["standard_price"]), "standard price (no option chosen yet)"


def _date(at: str | None) -> str:
    if not at:
        return ""
    try:
        return datetime.fromisoformat(str(at).replace("Z", "+00:00")).strftime("%-d %b %Y")
    except ValueError:
        return str(at)[:10]


# ---------------------------------------------------------------- open points

def open_points(tender: dict, totals: dict) -> list[dict]:
    """Export checklist (not blocking). Each item: {key, ok, title, detail, link}."""
    pts: list[dict] = []
    priceable = totals.get("lines_priceable", 0)
    evaluated = _evaluate(tender, _settings(tender))
    catalog_lines = sum(1 for p in evaluated if p["status"] == "calculated")
    no_price = [p for p in evaluated if p["status"] not in ("non_priceable", "out_of_scope")
                and not (p.get("pricing") or {}).get("direct") and not p.get("confirmed")]
    to_review = totals.get("to_review", 0)
    pts.append({"key": "lines", "ok": to_review == 0,
                "title": f"{catalog_lines} of {priceable} lines priced from the catalogue",
                "detail": (f"{to_review} line{'s' if to_review != 1 else ''} still to review" if to_review
                           else "No lines left to review"), "link": "lines"})

    n_budget = totals.get("budget_lines", 0)
    if n_budget:
        q = totals.get("quotes_entered", 0)
        pts.append({"key": "budget", "ok": False,
                    "title": f"{n_budget} line{'s' if n_budget != 1 else ''} on budget prices",
                    "detail": (f"CHF {P.swiss(totals['budget_direct'], 0)} of direct cost "
                               f"({P._half_up(totals['budget_share'] * 100)}%). "
                               + ("No supplier quotes entered yet." if not q else
                                  f"{q} supplier quote{'s' if q != 1 else ''} entered.")),
                    "link": "lines"})
    else:
        pts.append({"key": "budget", "ok": True, "title": "No lines on budget prices",
                    "detail": f"{totals.get('quotes_entered', 0)} supplier quotes entered", "link": "lines"})

    if no_price:
        n_sup = sum(1 for p in no_price if p["status"] == "supplier_quote_required")
        n_qty = sum(1 for p in no_price if p["status"] == "insufficient_information")
        n_other = len(no_price) - n_sup - n_qty
        bits = []
        if n_sup:
            bits.append(f"{n_sup} not in the catalogue (supplier requests are drafted)")
        if n_qty:
            bits.append(f"{n_qty} without a quantity in the tender")
        if n_other:
            bits.append(f"{n_other} without a product")
        pts.append({"key": "unpriced", "ok": False,
                    "title": f"{len(no_price)} line{'s' if len(no_price) != 1 else ''} without a price",
                    "detail": "; ".join(bits) + ". They are not in the direct cost yet.", "link": "lines"})

    ins = tender.get("instructions") or []
    if ins:
        open_ins = [i for i in ins if (i.get("state") or "open") == "open"]
        cost = sum(1 for i in open_ins if i.get("kind") == "cost")
        risk = sum(1 for i in open_ins if i.get("kind") == "risk")
        bits = []
        if cost:
            bits.append(f"{cost} add{'s' if cost == 1 else ''} cost")
        if risk:
            bits.append(f"{risk} {'is a risk' if risk == 1 else 'are risks'}")
        pts.append({"key": "instructions", "ok": not open_ins,
                    "title": (f"{len(open_ins)} of {len(ins)} instructions not confirmed" if open_ins
                              else f"All {len(ins)} instructions confirmed"),
                    "detail": ", ".join(bits), "link": "instructions"})

    detail = totals.get("project_costs_detail") or []
    on = [d for d in detail if d["on"]]
    off = [d for d in detail if not d["on"]]
    if not on:
        names = ", ".join(ref.PROJECT_COST_LABELS[d["key"]].split(" and ")[0].lower()
                          if d["key"] == "engineering" else ref.PROJECT_COST_LABELS[d["key"]].lower() for d in detail)
        pts.append({"key": "project_costs", "ok": False, "title": "Project costs not included",
                    "detail": f"{names} (CHF {P.swiss(totals.get('project_costs_all', 0), 0)} if all on)", "link": "pricing"})
    else:
        pts.append({"key": "project_costs", "ok": not off,
                    "title": f"Project costs CHF {P.swiss(totals['project_costs'], 0)} included",
                    "detail": ("Not included: " + ", ".join(d["label"] for d in off)) if off else "All project costs included",
                    "link": "pricing"})

    try:  # changes by people and Assist (and low-confidence AI picks) a senior has not approved yet
        from . import review
        rs = review.summary(tender)
    except Exception:
        rs = None
    if rs and (rs["changes"] or rs["ai_to_review"]):
        n, ai = rs["to_review"], rs["ai_to_review"]
        bits = [f"{rs['approved']} approved", f"{rs['reverted']} reverted"]
        if ai:
            bits.append(f"{ai} low-confidence AI pick{'s' if ai != 1 else ''} not reviewed")
        pts.append({"key": "reviews", "ok": not n and not ai,
                    "title": (f"{n} change{'s' if n != 1 else ''} not reviewed by a senior" if n
                              else f"All {rs['changes']} changes reviewed by a senior" if rs["changes"]
                              else "No changes by people yet"),
                    "detail": ", ".join(bits), "link": "changes"})

    if totals.get("offer"):
        o = totals["offer"]
        pts.append({"key": "offer", "ok": True, "title": f"Offer price CHF {P.swiss(o['bid'], 0)}",
                    "detail": f"{o['name']}, {o['margin_pct']:.1f}% margin · {P._half_up(o['win'] * 100)}% win chance",
                    "link": "pricing"})
    else:
        pts.append({"key": "offer", "ok": False, "title": "Offer price not chosen",
                    "detail": f"Standard price would be CHF {P.swiss(totals['standard_price'], 0)}", "link": "pricing"})
    return pts


# ---------------------------------------------------------------- settings with sources

def setting_rows(settings: dict) -> list[dict]:
    """Every pricing setting with value text, source and note (report and audit)."""
    s = ref.merge_settings(settings)
    d = ref.default_settings()
    changed_by = s.get("changed_by") or {}
    notes = s.get("notes") or {}

    def source(*keys):
        for k in keys:
            c = changed_by.get(k)
            if c:
                return f"Your call · {c.get('by', 'You')} · {_date(c.get('at'))}".rstrip(" ·")
        return "Your call"

    def word(table, v):
        for k, w in table.items():
            if abs(k - v) < 1e-9:
                return w
        return None

    rows = []

    def add(label, value, is_default, keys, note=""):
        rows.append({"label": label, "value": value,
                     "source": "Project parameters" if is_default else source(*keys), "note": note or "",
                     "changed": not is_default})

    pts = lambda v: f"{'+' if v > 0 else '-' if v < 0 else ''}{P._pts(v)} points"
    for k, table, extra in (("competition", COMPETITION_WORDS, ""), ("appetite", APPETITE_WORDS, ""),
                            ("sensitivity", SENSITIVITY_WORDS, " per 1% above standard")):
        w = word(table, s[k])
        val = (f"{w} ({pts(s[k]) if k != 'sensitivity' else '-' + str(P._pts(s[k])) + ' points'}{extra})" if w
               else f"{s[k]:+.2f}")
        add(ref.SETTING_LABELS[k], val, abs(s[k] - d[k]) < 1e-9, [k], notes.get(k, ""))
    add(ref.SETTING_LABELS["ref_win"], f"{P._half_up(s['ref_win'] * 100)}%", abs(s["ref_win"] - d["ref_win"]) < 1e-9, ["ref_win"])
    for name, m in s["margins"].items():
        add(f"{name} margin", f"{float(m):.1f}%", abs(float(m) - d["margins"].get(name, -1)) < 1e-9,
            [f"margins.{name}", name, "margins"])
    for k in ("overhead_pct", "transport_pct", "material_waste_pct", "small_material_pct", "supplier_contingency_pct"):
        add(ref.SETTING_LABELS[k], f"{s[k]:.1f}%", abs(s[k] - d[k]) < 1e-9, [k])
    for k in ("include_small_material", "include_supplier_contingency"):
        add(ref.SETTING_LABELS[k], "Yes" if s[k] else "No", bool(s[k]) == bool(d[k]), [k])
    for k, v in s["factors"].items():
        op = "÷" if k in ref.INVERTED_FACTORS else "×"
        add(f"Site factor: {ref.factor_label(k)}", f"{op} {float(v):.2f}", abs(float(v) - d["factors"].get(k, -1)) < 1e-9,
            [f"factors.{k}", k, "factors"])
    amounts = ref.project_cost_amounts()
    labels = ref.project_cost_labels()
    for k, a in amounts.items():
        on = bool((s.get("project_costs") or {}).get(k))
        add(f"Project cost: {labels[k]}", f"{'In the price' if on else 'Not included'} (CHF {P.swiss(a, 0)})",
            on == bool(d["project_costs"].get(k)), [f"project_costs.{k}", k, "project_costs"])
    return rows


# ---------------------------------------------------------------- xlsx

_HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
_HDR_FILL = PatternFill("solid", fgColor=INK)
_BODY = Font(name="Arial", size=10)
_BOLD = Font(name="Arial", size=10, bold=True)
_MUTED = Font(name="Arial", size=9, color=MUTED)
_TOP = Border(top=Side(style="thin", color=INK))
CHF = "#,##0.00"
CHF0 = "#,##0"


def _name(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9_]", "_", s)


def _header(ws, cols: list[tuple[str, float]], row: int = 1):
    for i, (label, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill = _HDR_FONT, _HDR_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _comment(cell, text: str):
    cell.comment = Comment(text, "ctrl-free", width=260, height=60)


def build_xlsx(tender: dict, totals: dict, settings: dict) -> bytes:
    """Offer workbook: Offer, Summary, Catalog, Install rules, Parameters (named cells)."""
    s = ref.merge_settings(settings)
    positions = _evaluate(tender, s)
    catalog = ref.load_catalog()
    rules = ref.load_rules()

    wb = Workbook()
    ws_offer = wb.active
    ws_offer.title = "Offer"
    ws_sum = wb.create_sheet("Summary")
    ws_cat = wb.create_sheet("Catalog")
    ws_rules = wb.create_sheet("Install rules")
    ws_par = wb.create_sheet("Parameters")

    # -- Parameters (named cells) ---------------------------------------
    _header(ws_par, [("Setting", 44), ("Value", 14), ("Name", 30), ("Source", 40)])
    defaults = ref.default_settings()
    names: dict[str, str] = {}
    prow = [2]

    def par(name: str, label: str, value, fmt: str | None = None, source: str | None = None):
        r = prow[0]
        ws_par.cell(row=r, column=1, value=label).font = _BODY
        c = ws_par.cell(row=r, column=2, value=value)
        c.font = _BOLD
        if fmt:
            c.number_format = fmt
        ws_par.cell(row=r, column=3, value=name).font = _MUTED
        ws_par.cell(row=r, column=4, value=source or "project_parameters.json").font = _MUTED
        wb.defined_names[name] = DefinedName(name, attr_text=f"Parameters!$B${r}")
        names[name] = f"Parameters!$B${r}"
        prow[0] += 1

    def src(changed: bool) -> str:
        return "Your call (this tender)" if changed else "project_parameters.json"

    par("waste_pct", "Material waste %", s["material_waste_pct"], "0.0", src(s["material_waste_pct"] != defaults["material_waste_pct"]))
    par("small_material_pct", "Small material % (of catalogue material)", s["small_material_pct"], "0.0")
    par("include_small_material", "Small material in the price (1 = yes)", int(bool(s["include_small_material"])), "0", "Pricing settings")
    par("supplier_contingency_pct", "Supplier package contingency % (of budget material)", s["supplier_contingency_pct"], "0.0")
    par("include_supplier_contingency", "Supplier contingency in the price (1 = yes)", int(bool(s["include_supplier_contingency"])), "0", "Pricing settings")
    par("transport_pct", "Transport and equipment %", s["transport_pct"], "0.0", src(s["transport_pct"] != defaults["transport_pct"]))
    par("overhead_pct", "Overhead %", s["overhead_pct"], "0.0", src(s["overhead_pct"] != defaults["overhead_pct"]))
    par("round_to", "Offer price rounded up to (CHF)", s["round_to"], CHF0)
    for k, v in s["factors"].items():
        par(f"f_{_name(k)}", f"Site factor: {ref.factor_label(k)}" + (" (hours ÷)" if k in ref.INVERTED_FACTORS else " (hours ×)"),
            float(v), "0.00", src(float(v) != defaults["factors"].get(k)))
    for n, m in s["margins"].items():
        par(f"margin_{_name(n)}", f"{n} margin %", float(m), "0.0", src(float(m) != defaults["margins"].get(n)))
    par("ref_win", "Win chance at the standard price", s["ref_win"], "0%", src(s["ref_win"] != defaults["ref_win"]))
    par("appetite", "How much we want this job (win points)", s["appetite"], "+0.00;-0.00;0.00", src(s["appetite"] != defaults["appetite"]))
    par("competition", "Competition (win points)", s["competition"], "+0.00;-0.00;0.00", src(s["competition"] != defaults["competition"]))
    par("sensitivity", "Price sensitivity (win points per 1% above standard)", s["sensitivity"], "0.00", src(s["sensitivity"] != defaults["sensitivity"]))
    par("min_win", "Lowest win chance", s["min_win"], "0%")
    par("max_win", "Highest win chance", s["max_win"], "0%")
    amounts = ref.project_cost_amounts()
    pc_labels = ref.project_cost_labels()
    for k, a in amounts.items():
        par(f"pc_{k}_amount", f"Project cost: {pc_labels[k]} (CHF)", a, CHF0)
        par(f"pc_{k}_on", f"Project cost: {pc_labels[k]} in the price (1 = yes)", int(bool(s["project_costs"].get(k))), "0", "Pricing settings")

    # -- Catalog ---------------------------------------------------------
    cat_cols = ["product_id", "category", "subcategory", "manufacturer", "model", "description", "unit",
                "list_price_chf", "discount_pct", "net_cost_chf", "install_code", "source_type", "supplier_quote_required"]
    _header(ws_cat, [(c, 16 if c != "description" else 44) for c in cat_cols])
    for i, row in enumerate(catalog.values(), start=2):
        for j, c in enumerate(cat_cols, start=1):
            v = row.get(c)
            ws_cat.cell(row=i, column=j, value=v if not isinstance(v, bool) else str(v).lower()).font = _BODY
    cat_last = len(catalog) + 1
    net_col = get_column_letter(cat_cols.index("net_cost_chf") + 1)

    # -- Install rules ---------------------------------------------------
    rule_cols = ["install_code", "work_type", "unit", "base_hours_per_unit", "labour_role", "default_hourly_rate_chf", "factor_keys"]
    _header(ws_rules, [(c, 18 if c not in ("work_type", "factor_keys") else 44) for c in rule_cols])
    for i, row in enumerate(rules.values(), start=2):
        for j, c in enumerate(rule_cols, start=1):
            v = row.get(c)
            ws_rules.cell(row=i, column=j, value="|".join(v) if isinstance(v, list) else v).font = _BODY
    rules_last = len(rules) + 1
    bh_col = get_column_letter(rule_cols.index("base_hours_per_unit") + 1)
    rate_col = get_column_letter(rule_cols.index("default_hourly_rate_chf") + 1)

    # -- Offer -------------------------------------------------------------
    cols = [("LV id", 9), ("System", 26), ("Section", 16), ("Page", 6), ("Tender text", 48), ("Qty", 9), ("Unit", 6),
            ("Product id", 16), ("Product description", 34), ("Net unit price", 12), ("Price source", 15),
            ("Waste %", 8), ("Material", 13), ("Install rule", 16), ("Base h/unit", 10), ("Factor product", 10),
            ("Hours", 10), ("Rate", 8), ("Labour", 12), ("Direct", 13), ("Status", 22), ("Status reason", 40),
            ("Source page", 10)]
    _header(ws_offer, cols)
    ws_offer.row_dimensions[1].height = 30
    C = {label: get_column_letter(i) for i, (label, _) in enumerate(cols, start=1)}
    first = 2
    attention: list[int] = []
    for i, p in enumerate(positions, start=first):
        pr, st = p["pricing"], p["status"]
        page = _page(p)
        sec = [x for x in (p.get("section"), p.get("section_name")) if x]
        vals = {"LV id": p.get("lv_id"), "System": p.get("system"),
                "Section": " ".join(dict.fromkeys(sec)) or None, "Page": page, "Tender text": _title(p), "Status": P.STATUS_LABELS.get(st, (st,))[0] + f" ({st})",
                "Status reason": p.get("status_reason"), "Source page": f"p. {page}" if page else None}
        for label, v in vals.items():
            ws_offer[f"{C[label]}{i}"] = v
        q = P.effective_quantity(p)
        ws_offer[f"{C['Qty']}{i}"] = q
        ws_offer[f"{C['Unit']}{i}"] = p.get("unit") or (pr or {}).get("unit")
        pid = P.effective_product_id(p)
        ws_offer[f"{C['Product id']}{i}"] = pid
        prod = catalog.get(pid) if pid else None
        ws_offer[f"{C['Product description']}{i}"] = (prod or {}).get("description")
        if page:
            _comment(ws_offer[f"{C['Page']}{i}"], f"Tender p. {page}")
        if prod:
            _comment(ws_offer[f"{C['Product id']}{i}"], f"catalog.csv {pid}")
        if q is not None and p.get("overrides", {}).get("quantity") is not None:
            _comment(ws_offer[f"{C['Qty']}{i}"], f"Changed by you; tender p. {page} says {P.fmt_qty(p.get('quantity'))}")
        elif q is not None and page:
            _comment(ws_offer[f"{C['Qty']}{i}"], f"Tender p. {page}")
        if p.get("status_reason"):
            _comment(ws_offer[f"{C['Status']}{i}"], p["status_reason"])
        if p["needs_review"]:
            attention.append(i)

        if not P.included(st, pr):
            continue  # status shown, amounts blank
        r = str(i)
        # unit price: formula into Catalog for catalogue / budget prices, value for quotes and manual prices
        if pr["unit_price_source"] in ("catalog", "supplier_package") and prod and not pr.get("price_list"):
            ws_offer[C["Net unit price"] + r] = f"=INDEX(Catalog!${net_col}$2:${net_col}${cat_last},MATCH({C['Product id']}{r},Catalog!$A$2:$A${cat_last},0))"
            _comment(ws_offer[C["Net unit price"] + r], f"catalog.csv {pid} net_cost_chf" + (" (budget price)" if pr["unit_price_source"] == "supplier_package" else ""))
        else:
            ws_offer[C["Net unit price"] + r] = pr["net_unit_cost"]
            _comment(ws_offer[C["Net unit price"] + r], pr["sources"]["unit_price"])
        ws_offer[C["Price source"] + r] = pr["unit_price_source"]
        ws_offer[C["Waste %"] + r] = "=waste_pct"
        ws_offer[C["Material"] + r] = f"=ROUND({C['Qty']}{r}*{C['Net unit price']}{r}*(1+{C['Waste %']}{r}/100),2)"
        code = pr["install_code"]
        rule = rules.get(code) if code else None
        ws_offer[C["Install rule"] + r] = code
        if rule:
            _comment(ws_offer[C["Install rule"] + r], f"productivity.csv {code}: {rule['work_type']}")
            match = f"MATCH({C['Install rule']}{r},'Install rules'!$A$2:$A${rules_last},0)"
            ws_offer[C["Base h/unit"] + r] = f"=INDEX('Install rules'!${bh_col}$2:${bh_col}${rules_last},{match})"
            ws_offer[C["Rate"] + r] = f"=INDEX('Install rules'!${rate_col}$2:${rate_col}${rules_last},{match})"
            terms = [(f"1/f_{_name(t['key'])}" if t["op"] == "÷" else f"f_{_name(t['key'])}") for t in pr["factors"]]
            ws_offer[C["Factor product"] + r] = "=" + ("*".join(terms) if terms else "1")
            _comment(ws_offer[C["Factor product"] + r], " ".join(f"{t['op']} {t['value']:.2f} {t['label']}" for t in pr["factors"]) or "No site factors")
        elif pr["unit_price_source"] == "labour":
            # hourly work: the tender quantity is the hours, priced at the role's hourly rate
            ws_offer[C["Base h/unit"] + r] = 1
            ws_offer[C["Rate"] + r] = pr["rate"]
            _comment(ws_offer[C["Rate"] + r], f"Hourly work: productivity.csv rate for {pr['role'].replace('_', ' ')}")
            ws_offer[C["Factor product"] + r] = 1
        else:
            ws_offer[C["Base h/unit"] + r] = 0
            ws_offer[C["Rate"] + r] = 0
            ws_offer[C["Factor product"] + r] = 1
        if pr.get("hours_overridden"):
            ws_offer[C["Hours"] + r] = pr["hours"]
            _comment(ws_offer[C["Hours"] + r], f"Set by you; the install rule gives {P.swiss(pr['calculated_hours'])} h")
        else:
            ws_offer[C["Hours"] + r] = f"={C['Qty']}{r}*{C['Base h/unit']}{r}*{C['Factor product']}{r}"
        ws_offer[C["Labour"] + r] = f"=ROUND({C['Hours']}{r}*{C['Rate']}{r},2)"
        ws_offer[C["Direct"] + r] = f"={C['Material']}{r}+{C['Labour']}{r}"

    last = max(first, len(positions) + 1)
    for row in ws_offer.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.font = _BODY
    for i in attention:
        ws_offer[f"{C['Status']}{i}"].font = Font(name="Arial", size=10, color=ATTN)
    for label, fmt in (("Net unit price", CHF), ("Material", CHF), ("Labour", CHF), ("Direct", CHF), ("Hours", "0.00"),
                       ("Base h/unit", "0.00"), ("Factor product", "0.0000"), ("Rate", CHF0), ("Qty", "#,##0.###")):
        for i in range(first, last + 1):
            ws_offer[f"{C[label]}{i}"].number_format = fmt
    for i in range(first, last + 1):
        ws_offer[f"{C['Tender text']}{i}"].alignment = Alignment(wrap_text=False)
    tot = last + 1
    ws_offer[f"{C['Tender text']}{tot}"] = "Total"
    for label in ("Material", "Labour", "Direct"):
        ws_offer[f"{C[label]}{tot}"] = f"=SUM({C[label]}{first}:{C[label]}{last})"
        ws_offer[f"{C[label]}{tot}"].number_format = CHF
    for col in range(1, len(cols) + 1):
        ws_offer.cell(row=tot, column=col).font = _BOLD
        ws_offer.cell(row=tot, column=col).border = _TOP
    ws_offer.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    # -- Summary -----------------------------------------------------------
    ws_sum.column_dimensions["A"].width = 48
    for col in "BCDEF":
        ws_sum.column_dimensions[col].width = 16
    ws_sum["A1"] = tender.get("name") or "Offer"
    ws_sum["A1"].font = Font(name="Arial", size=14, bold=True)
    ws_sum["A2"] = " · ".join(x for x in (tender.get("client"), tender.get("trade"),
                                         f"generated {datetime.now():%d.%m.%Y %H:%M}") if x)
    ws_sum["A2"].font = _MUTED
    _hdr_row = 4
    for j, h in enumerate(("How the cost adds up", "CHF (formula)", "ctrl-free value"), start=1):
        c = ws_sum.cell(row=_hdr_row, column=j, value=h)
        c.font, c.fill = _HDR_FONT, _HDR_FILL
    rng = lambda label: f"Offer!${C[label]}${first}:${C[label]}${last}"
    pc_formula = "+".join(f"pc_{k}_on*pc_{k}_amount" for k in amounts) or "0"
    lines = [
        ("material", "Material", f"=SUM({rng('Material')})", totals.get("material")),
        ("labour", "Labour", f"=SUM({rng('Labour')})", totals.get("labour")),
        ("direct", f"Direct cost, {totals.get('lines_priced', 0)} lines", f"=SUM({rng('Direct')})", totals.get("direct")),
        ("project_costs", "+ Project costs switched on", f"={pc_formula}", totals.get("project_costs")),
        ("small_material", "+ Small material (of catalogue material)",
         f"=IF(include_small_material=1,small_material_pct/100*SUMIF({rng('Price source')},\"catalog\",{rng('Material')}),0)",
         totals.get("small_material")),
        ("supplier_contingency", "+ Supplier package contingency (of budget material)",
         f"=IF(include_supplier_contingency=1,supplier_contingency_pct/100*SUMIF({rng('Price source')},\"supplier_package\",{rng('Material')}),0)",
         totals.get("supplier_contingency")),
    ]
    R: dict[str, str] = {}
    r = _hdr_row + 1
    for key, label, formula, val in lines:
        ws_sum[f"A{r}"], ws_sum[f"B{r}"], ws_sum[f"C{r}"] = label, formula, val
        R[key] = f"B{r}"
        r += 1
    rest = [
        ("base", "Base for transport", f"={R['direct']}+{R['project_costs']}+{R['small_material']}+{R['supplier_contingency']}", None),
        ("transport", "+ Transport and equipment", "={base}*transport_pct/100", totals.get("transport")),
        ("overhead", "+ Overhead", "=({base}+{transport})*overhead_pct/100", totals.get("overhead")),
        ("cost", "Cost", "={base}+{transport}+{overhead}", totals.get("cost")),
        ("standard", "Standard price (standard margin, rounded up)", "=CEILING({cost}/(1-margin_Standard/100),round_to)",
         totals.get("standard_price")),
    ]
    for key, label, formula, val in rest:
        R[key] = f"B{r}"
        ws_sum[f"A{r}"], ws_sum[f"B{r}"], ws_sum[f"C{r}"] = label, formula.format(**R), val
        if key in ("cost", "standard"):
            ws_sum[f"A{r}"].font = ws_sum[f"B{r}"].font = _BOLD
            ws_sum[f"A{r}"].border = ws_sum[f"B{r}"].border = ws_sum[f"C{r}"].border = _TOP
        r += 1
    for rr in range(_hdr_row + 1, r):
        ws_sum[f"B{rr}"].number_format = ws_sum[f"C{rr}"].number_format = CHF
    wb.defined_names["total_direct"] = DefinedName("total_direct", attr_text=f"Summary!${R['direct'][0]}${R['direct'][1:]}")
    wb.defined_names["total_cost"] = DefinedName("total_cost", attr_text=f"Summary!${R['cost'][0]}${R['cost'][1:]}")

    r += 1
    ws_sum[f"A{r}"] = "Pick a price"
    ws_sum[f"A{r}"].font = Font(name="Arial", size=12, bold=True)
    r += 1
    for j, h in enumerate(("Option", "Margin %", "Offer price", "Win chance", "Profit if won", "Expected profit"), start=1):
        c = ws_sum.cell(row=r, column=j, value=h)
        c.font, c.fill = _HDR_FONT, _HDR_FILL
    r += 1
    std_ref = f"({R['cost']}/(1-margin_Standard/100))"
    chosen_row = None
    for sc in totals.get("scenarios") or []:
        n = _name(sc["name"])
        ws_sum[f"A{r}"] = sc["name"] + ("  · best expected profit" if sc.get("best") else "") + \
            ("  · your offer" if totals.get("chosen") == sc["name"] else "")
        ws_sum[f"B{r}"] = f"=margin_{n}"
        ws_sum[f"C{r}"] = f"=CEILING({R['cost']}/(1-B{r}/100),round_to)"
        ws_sum[f"D{r}"] = f"=MIN(max_win,MAX(min_win,ref_win+appetite+competition-sensitivity*(C{r}/{std_ref}-1)*100))"
        ws_sum[f"E{r}"] = f"=C{r}-{R['cost']}"
        ws_sum[f"F{r}"] = f"=D{r}*E{r}"
        ws_sum[f"B{r}"].number_format = "0.0"
        ws_sum[f"C{r}"].number_format = CHF0
        ws_sum[f"D{r}"].number_format = "0%"
        ws_sum[f"E{r}"].number_format = ws_sum[f"F{r}"].number_format = CHF
        if totals.get("chosen") == sc["name"]:
            chosen_row = r
            for col in "ABCDEF":
                ws_sum[f"{col}{r}"].fill = PatternFill("solid", fgColor="EEF2FD")
                ws_sum[f"{col}{r}"].font = _BOLD
        r += 1
    r += 1
    ws_sum[f"A{r}"] = "Offer price"
    ws_sum[f"A{r}"].font = Font(name="Arial", size=12, bold=True)
    if chosen_row:
        ws_sum[f"B{r}"] = f"=C{chosen_row}"
        ws_sum[f"B{r}"].number_format = CHF0
        ws_sum[f"B{r}"].font = Font(name="Arial", size=12, bold=True)
        ws_sum[f"C{r}"] = totals.get("offer_price")
        ws_sum[f"C{r}"].number_format = CHF0
    else:
        ws_sum[f"B{r}"] = "Not chosen yet"
    r += 2
    ws_sum[f"A{r}"] = totals.get("win_explain")
    ws_sum[f"A{r}"].font = _MUTED
    r += 1
    ws_sum[f"A{r}"] = ("Column B recalculates from the Offer and Parameters sheets; column C is the value ctrl-free "
                       "computed when the file was made. Win chance = win at the standard price + appetite + "
                       "competition - sensitivity × % above the standard price.")
    ws_sum[f"A{r}"].font = _MUTED

    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF report

def _pdf_text(s) -> str:
    """Helvetica (WinAnsi) safe text, XML-escaped for Paragraph."""
    s = "" if s is None else str(s)
    s = s.replace("−", "-").replace("≥", ">=").replace("≤", "<=").replace("→", "->").replace(" ", " ")
    s = s.encode("cp1252", errors="replace").decode("cp1252")
    return escape(s)


def build_report_pdf(tender: dict, totals: dict, settings: dict, open_points: list[dict]) -> bytes:
    """Pricing report (A4): offer price, build-up, every setting with its source, open points,
    systems, positions, assumptions and the status legend."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    s = ref.merge_settings(settings)
    positions = _evaluate(tender, s)
    ink, muted, rule, accent, attn = (colors.HexColor("#" + x) for x in (INK, MUTED, RULE, ACCENT, ATTN))
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=9, leading=12, textColor=ink)
    small = ParagraphStyle("s", parent=base, fontSize=7.5, leading=9.5)
    small_r = ParagraphStyle("sr", parent=small, alignment=TA_RIGHT)
    mute = ParagraphStyle("m", parent=base, textColor=muted, fontSize=8, leading=10.5)
    h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=16, leading=20, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=10, spaceAfter=4)
    big = ParagraphStyle("big", parent=base, fontName="Helvetica-Bold", fontSize=20, leading=24)
    para = lambda t, st=base: Paragraph(_pdf_text(t), st)

    doc_w = A4[0] - 30 * mm

    def table(rows, widths, header=True, right_cols=(), bold_rows=(), font=8):
        t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
        style = [("FONT", (0, 0), (-1, -1), "Helvetica", font), ("TEXTCOLOR", (0, 0), (-1, -1), ink),
                 ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.3, rule),
                 ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                 ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]
        if header:
            style += [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", font), ("LINEBELOW", (0, 0), (-1, 0), 0.8, ink)]
        for c in right_cols:
            style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
        for r in bold_rows:
            style += [("FONT", (0, r), (-1, r), "Helvetica-Bold", font), ("LINEABOVE", (0, r), (-1, r), 0.8, ink)]
        t.setStyle(TableStyle(style))
        return t

    story = []
    story.append(para("Pricing report", h1))
    meta = " · ".join(x for x in (tender.get("name"), tender.get("client"), tender.get("trade")) if x)
    story.append(para(meta, mute))
    due = tender.get("due_at")
    story.append(para(" · ".join(x for x in (
        f"Due {datetime.fromisoformat(due).strftime('%d.%m.%Y %H:%M')}" if due else None,
        f"File {tender.get('file_name')}" if tender.get("file_name") else None,
        f"Generated {datetime.now():%d.%m.%Y %H:%M}") if x), mute))
    story.append(Spacer(1, 8))

    # offer price
    if totals.get("offer"):
        o = totals["offer"]
        story.append(para(f"Offer price CHF {P.swiss(o['bid'], 0)}", big))
        story.append(para(f"{o['name']} · {o['margin_pct']:.1f}% margin · {P._half_up(o['win'] * 100)}% win chance · "
                          f"CHF {P.swiss(o['profit'], 0)} profit if won"))
    else:
        story.append(para("Offer price not chosen yet", big))
        story.append(para(f"Standard price (standard margin {totals.get('standard_margin_pct', 0):.1f}%): "
                          f"CHF {P.swiss(totals['standard_price'], 0)}"))
    story.append(para(totals.get("win_explain"), mute))

    story.append(para("Pick a price", h2))
    rows = [["Option", "Margin", "Offer price", "Win chance", "Profit if won", "Expected profit", ""]]
    for sc in totals.get("scenarios") or []:
        tags = []
        if sc.get("best"):
            tags.append("best expected profit")
        if totals.get("chosen") == sc["name"]:
            tags.append("your offer")
        rows.append([sc["name"], f"{sc['margin_pct']:.1f}%", P.swiss(sc["bid"], 0), f"{P._half_up(sc['win'] * 100)}%",
                     P.swiss(sc["profit"], 0), P.swiss(sc["expected"], 0), ", ".join(tags)])
    story.append(table(rows, [28 * mm, 18 * mm, 26 * mm, 22 * mm, 26 * mm, 28 * mm, doc_w - 148 * mm], right_cols=(1, 2, 3, 4, 5)))
    story.append(para("Expected profit = win chance × profit if won.", mute))

    story.append(para("How the cost adds up", h2))
    rows = [[b["label"], P.swiss(b["value"], 0 if b["key"] == "offer" else 2)] for b in totals.get("buildup") or []]
    bold = [i for i, b in enumerate(totals.get("buildup") or []) if b.get("strong")]
    story.append(table(rows, [doc_w - 40 * mm, 40 * mm], header=False, right_cols=(1,), bold_rows=bold, font=9))

    story.append(para("Open points before you submit", h2))
    rows = [["", "Point", "Detail"]]
    for pt in open_points or []:
        rows.append(["OK" if pt.get("ok") else "Open", para(pt.get("title"), small), para(pt.get("detail"), small)])
    t = table(rows, [12 * mm, 70 * mm, doc_w - 82 * mm])
    for i, pt in enumerate(open_points or [], start=1):
        t.setStyle(TableStyle([("TEXTCOLOR", (0, i), (0, i), muted if pt.get("ok") else attn)]))
    story.append(t)

    story.append(para("Settings and where they come from", h2))
    rows = [["Setting", "Value", "Source", "Note"]]
    srows = setting_rows(s)
    for sr in srows:
        rows.append([para(sr["label"], small), para(sr["value"], small), para(sr["source"], small), para(sr["note"], small)])
    t = table(rows, [62 * mm, 40 * mm, 38 * mm, doc_w - 140 * mm])
    for i, sr in enumerate(srows, start=1):
        if sr["changed"]:
            t.setStyle(TableStyle([("BACKGROUND", (0, i), (-1, i), colors.HexColor("#EEF2FD"))]))
    story.append(t)

    if totals.get("by_system"):
        story.append(para("By system", h2))
        rows = [["System", "Lines", "Priced", "To review", "Material", "Labour", "Direct"]]
        for r_ in totals["by_system"]:
            rows.append([para(r_["system"], small), r_["lines"], r_["priced"], r_["to_review"], P.swiss(r_["material"]),
                         P.swiss(r_["labour"]), P.swiss(r_["direct"])])
        rows.append(["Total", "", "", "", P.swiss(totals["material"]), P.swiss(totals["labour"]), P.swiss(totals["direct"])])
        story.append(table(rows, [doc_w - 128 * mm, 14 * mm, 14 * mm, 16 * mm, 28 * mm, 28 * mm, 28 * mm],
                           right_cols=(1, 2, 3, 4, 5, 6), bold_rows=(len(rows) - 1,)))

    story.append(para("Lines", h2))
    rows = [["LV", "p.", "Tender text", "Qty", "Product", "Unit price", "Direct", "Status"]]
    for p in positions:
        pr = p["pricing"]
        inc = P.included(p["status"], pr)
        q = P.effective_quantity(p)
        rows.append([p.get("lv_id"), _page(p) or "", para(_title(p)[:90], small),
                     para(f"{P.fmt_qty(q)} {p.get('unit') or ''}".strip(), small_r),
                     para(P.effective_product_id(p) or "", small),
                     P.swiss(pr["net_unit_cost"]) if inc else "", P.swiss(pr["material"] + pr["labour"]) if inc else "",
                     para(P.STATUS_LABELS.get(p["status"], (p["status"],))[0], small)])
    story.append(table(rows, [14 * mm, 8 * mm, doc_w - 128 * mm, 16 * mm, 26 * mm, 20 * mm, 22 * mm, 22 * mm],
                       right_cols=(1, 5, 6), font=7.5))

    story.append(para("Assumptions", h2))
    n_budget = totals.get("budget_lines", 0)
    assumptions = [
        "Every amount is computed from the tender quantities and the reference data (catalogue, install rules, "
        "project parameters). Products are suggested by matching; prices are never produced by the language model.",
        f"Material = quantity × net unit cost × (1 + {P.fmt_pct(s['material_waste_pct'])} waste).",
        "Hours = quantity × base hours of the install rule × site factors; site efficiency divides, the other "
        "factors multiply. Sports hall systems use the sports hall ceiling height factor.",
        "Labour = hours × the hourly rate of the install rule's role.",
        "Lines the tender counts in hours with no product (hourly work, Regie) are our own labour: the tender's "
        "hours × the hourly rate of the matching role (installer unless the line names commissioning), no material.",
        f"{n_budget} line{'s' if n_budget != 1 else ''} use supplier package budget prices "
        f"(CHF {P.swiss(totals.get('budget_direct', 0), 0)} direct) until a supplier quote is entered.",
        "Small material, when switched on, applies to catalogue material only; supplier package contingency applies "
        "to material on budget prices.",
        f"Transport and equipment ({s['transport_pct']:.1f}%) applies to direct cost plus project costs and allowances; "
        f"overhead ({s['overhead_pct']:.1f}%) applies to that plus transport.",
        f"Offer prices are cost ÷ (1 - margin), rounded up to CHF {P.swiss(s['round_to'], 0)}.",
        "Win chance = win chance at the standard price + appetite + competition - sensitivity × % above the "
        f"standard price, kept between {P._half_up(s['min_win'] * 100)}% and {P._half_up(s['max_win'] * 100)}%.",
        "Lines that are out of scope or not priceable are listed but not in the price.",
    ]
    for a in assumptions:
        story.append(para("- " + a, small))

    story.append(KeepTogether([para("Statuses", h2)] + [
        table([["Status", "Meaning"]] + [[para(f"{P.STATUS_LABELS[k][0]} ({k})", small), para(v, small)]
                                         for k, v in STATUS_HELP.items()], [55 * mm, doc_w - 55 * mm])]))

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(muted)
        canvas.drawString(15 * mm, 10 * mm, _pdf_text(f"ctrl-free · {tender.get('name') or ''}")[:110].replace("&amp;", "&"))
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm,
                            bottomMargin=16 * mm, title=f"Pricing report - {tender.get('name') or ''}",
                            author="ctrl-free")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# ---------------------------------------------------------------- SIA-like file

def build_sia_file(tender: dict, totals: dict) -> bytes:
    """Simplified SIA-451-like interchange file (semicolon CSV, UTF-8).

    Unit prices are selling prices: the offer price (chosen option, else the standard
    price) is spread over the priced lines in proportion to their direct cost.
    """
    s = _settings(tender)
    positions = _evaluate(tender, s)
    offer, how = _offer(totals)
    direct = float(totals.get("direct") or 0)
    k = offer / direct if direct else 0.0
    now = datetime.now()
    out = io.StringIO()
    w = lambda line: out.write(line + "\r\n")
    for line in (
        "# ctrl-free simplified LV interchange file (SIA-451-like)",
        "# This is a simplified, human-readable interchange file, NOT the official SIA 451 format.",
        "# Producing a real SIA 451 file (fixed records per the CRB/SIA specification) is the production next step.",
        f"# Tender: {tender.get('name') or ''}",
        f"# Client: {tender.get('client') or ''}",
        f"# Trade: {tender.get('trade') or ''}",
        f"# Source file: {tender.get('file_name') or ''}",
        f"# Created: {now:%Y-%m-%d %H:%M}",
        f"# Offer price CHF {offer:.2f} ({how}); direct cost CHF {direct:.2f}; markup factor {k:.6f}",
        "# Unit prices include transport, overhead, project costs and margin, spread over lines in proportion to direct cost.",
    ):
        w(line)
    w("Pos;LV;System;Kapitel;Seite;Text;Menge;Einheit;Einheitspreis CHF;Betrag CHF;Status")

    def cell(v) -> str:
        t = "" if v is None else str(v)
        t = t.replace("\r", " ").replace("\n", " ").strip()
        return f'"{t.replace(chr(34), chr(34) * 2)}"' if (";" in t or '"' in t) else t

    total = 0.0
    n = 0
    for p in positions:
        if (p.get("kind") or "item") in ("subtotal", "text_only"):
            continue
        n += 1
        pr = p["pricing"]
        q = P.effective_quantity(p)
        unit = SIA_UNITS.get(p.get("unit") or (pr or {}).get("unit") or "", p.get("unit_raw") or p.get("unit") or "")
        if P.included(p["status"], pr) and q:
            unit_price = round((pr["material"] + pr["labour"]) * k / q, 2)
            amount = round(q * unit_price, 2)
            total += amount
            up, am = f"{unit_price:.2f}", f"{amount:.2f}"
        else:
            up, am = "", ""
        w(";".join(cell(x) for x in (
            f"{n:04d}", p.get("lv_id"), p.get("system_key") or p.get("system"), p.get("section"), _page(p),
            _title(p), f"{q:g}" if q is not None else "", unit, up, am, p["status"])))
    w(f"TOTAL;;;;;Summe Positionen;;;;{total:.2f};")
    w(f"OFFER;;;;;Angebotspreis (gerundet);;;;{offer:.2f};")
    w(f"DIFF;;;;;Rundungsdifferenz;;;;{offer - total:.2f};")
    return out.getvalue().encode("utf-8")
