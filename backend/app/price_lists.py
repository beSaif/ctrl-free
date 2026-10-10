"""Supplier and negotiated price lists: net prices per catalogue product, with one clear precedence.

Equans buys from suppliers who each send their own catalogue, and negotiates its own net prices
on top. A price list is reference data next to catalog.csv:

    {id, name, supplier, kind: negotiated | supplier, valid_from, valid_to?, currency, file, active,
     rows: [{product_id | manufacturer + model (+ supplier article no.), net_price | list_price - discount, unit?}]}

Which unit price a line uses (the first that applies wins):
  1. a unit price the estimator typed on the line            (pricing.py: overrides.unit_price)
  2. the supplier's quote for this tender                     (pricing.py: supplier_quote)
  3. the active price lists, in priority order. By default Equans' negotiated lists come first,
     then supplier lists, the most recent valid_from first; the estimator can reorder them on
     the Reference page. A list counts only while it is active and valid today.
  4. the standard catalogue (catalog.csv net_cost_chf)
The order is global: it applies to every tender, and changing it re-prices them all (each line
whose unit price moves gets a log entry with the old and the new price).

Rows are matched to catalogue products by product_id, else by manufacturer + model (the list's
supplier stands in for a missing manufacturer), else by a model that is unique in the catalogue.
Rows that do not match, or whose unit differs from the catalogue unit, are kept and shown as
unmatched; they never price anything.

Storage (config.PRICE_LISTS_DIR): index.json (the lists, in priority order), <id>.json (parsed
rows and the column mapping), <id>.<ext> (the original upload).
"""
from __future__ import annotations

import contextlib
import contextvars
import csv
import io
import json
import re
import secrets
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from app import config
from app import reference as ref
from app.units import normalize_unit

KINDS = ("negotiated", "supplier")
KIND_LABELS = {"negotiated": "Equans negotiated prices", "supplier": "Supplier price list"}
PRECEDENCE = [
    "Unit price typed by the estimator on the line",
    "Supplier quote for this tender",
    "Active price lists, top to bottom (Equans negotiated lists first by default, then supplier lists, newest first)",
    "Standard catalogue (catalog.csv)",
]

# Columns a price list can have, and the headers we recognise for each (compared lower-case,
# umlauts folded, without spaces and punctuation).
FIELDS = ["product_id", "article_no", "manufacturer", "model", "description", "unit",
          "list_price", "discount_pct", "net_price"]
FIELD_LABELS = {
    "product_id": "Our product id", "article_no": "Supplier article no.", "manufacturer": "Manufacturer",
    "model": "Model / type", "description": "Description", "unit": "Unit", "list_price": "List price",
    "discount_pct": "Discount %", "net_price": "Net price",
}
ALIASES = {
    "product_id": ["productid", "produktid", "product", "produkt", "equansnr", "equansid", "equansartikelnr",
                   "interneartikelnr", "internalid", "ourid", "katalognrequans"],
    "article_no": ["artikelnr", "artnr", "artikelnummer", "articleno", "articlenumber", "articlenr", "sku",
                   "lieferantenartikelnr", "bestellnr", "bestellnummer", "supplierarticleno", "artikel"],
    "manufacturer": ["hersteller", "manufacturer", "make", "marke", "brand", "fabrikat"],
    "model": ["typ", "type", "modell", "model", "typenbezeichnung", "typbezeichnung", "modelnr"],
    "description": ["bezeichnung", "beschreibung", "description", "artikelbezeichnung", "produktbezeichnung",
                    "kurztext", "text", "name"],
    "unit": ["einheit", "me", "unit", "mengeneinheit", "uom", "masseinheit"],
    "list_price": ["listenpreis", "bruttopreis", "brutto", "listprice", "listpricechf", "grossprice", "gross",
                   "katalogpreis", "preisbrutto"],
    "discount_pct": ["rabatt", "rabattprozent", "rabattsatz", "discount", "discountpct"],
    "net_price": ["netto", "nettopreis", "nettochf", "netprice", "net", "netcost", "netcostchf", "preisnetto",
                  "einkaufspreis", "ekpreis", "preis", "price", "unitprice"],
}
MAX_ROWS = 20000


class _NoList(KeyError):
    pass


def _fold(s) -> str:
    s = str(s or "").lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]", "", s)


def _today() -> date:
    return date.today()


def _date(v) -> date | None:
    if not v:
        return None
    if isinstance(v, (date, datetime)):
        return v if isinstance(v, date) and not isinstance(v, datetime) else v.date()
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$", s)
    if m:
        return date(int(m[3]), int(m[2]), int(m[1]))
    raise ValueError(f"'{s}' is not a date (use YYYY-MM-DD)")


def fmt_date(v) -> str:
    """2026-01-01 -> "1 Jan 2026" (as in the log and the source sentence)."""
    d = _date(v)
    return f"{d.day} {d:%b %Y}" if d else ""


# ---------------------------------------------------------------- reading files

def _num(v, pct_cell: bool = False) -> float | None:
    """Prices as suppliers write them: 1'234.50, 1.234,50, 74,80, CHF 74.80, 12 %."""
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return round(float(v) * (100 if pct_cell else 1), 6)
    s = re.sub(r"(?i)chf|fr\.|sfr\.?|%|\s", "", str(v)).replace("’", "'").replace("'", "").replace("`", "")
    if not s or s in ("-", "–"):
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif s.count(",") == 1:
        s = s.replace(",", ".")
    elif s.count(",") > 1:
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def read_table(data: bytes, filename: str) -> list[list]:
    """All rows of the first sheet (xlsx) or the CSV, as lists of cell values.

    xlsx percentage cells come back as percent numbers (0.12 formatted 12% -> 12.0)."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # not a readable workbook
            raise ValueError(f"Could not read the Excel file: {e}") from None
        ws = wb.worksheets[0]
        rows = []
        for r in ws.iter_rows():
            rows.append([(round(c.value * 100, 6) if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
                          and "%" in str(getattr(c, "number_format", "") or "") else c.value) for c in r])
            if len(rows) > MAX_ROWS + 20:
                break
        wb.close()
        return rows
    if name.endswith(".xls"):
        raise ValueError("Old .xls files are not supported: save the list as .xlsx or .csv")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Could not read the file: use CSV (UTF-8 or Windows) or XLSX")
    sample = "\n".join(text.splitlines()[:20])
    counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
    delim = max(counts, key=counts.get) if any(counts.values()) else ","
    return [r for r in csv.reader(io.StringIO(text), delimiter=delim)]


def _field_for(header: str, exact_only: bool = False) -> str | None:
    h = _fold(header)
    if not h:
        return None
    for f in FIELDS:
        if h in ALIASES[f]:
            return f
    if exact_only:
        return None
    # "Bruttopreis CHF", "Netto-Preis (CHF)", "Rabatt in %": an alias at the start of the header
    best = None
    for f in FIELDS:
        for a in ALIASES[f]:
            if len(a) >= 4 and h.startswith(a) and (best is None or len(a) > best[1]):
                best = (f, len(a))
    return best[0] if best else None


def detect_mapping(headers: list[str]) -> dict[str, str | None]:
    """{field: header} from the file's own headers (German or English)."""
    out: dict[str, str | None] = {f: None for f in FIELDS}
    for exact in (True, False):
        for h in headers:
            if h in out.values():
                continue
            f = _field_for(h, exact_only=exact)
            if f and out[f] is None:
                out[f] = h
    return out


def _header_index(rows: list[list]) -> int:
    """The first row (of the first 15) that looks like a header: 2+ cells we recognise."""
    for i, r in enumerate(rows[:15]):
        if sum(1 for c in r if isinstance(c, str) and _field_for(c)) >= 2:
            return i
    return next((i for i, r in enumerate(rows) if any(_cell_text(c) for c in r)), 0)


def parse_table(data: bytes, filename: str, mapping: dict | None = None) -> dict:
    """{headers, mapping, title_lines, rows: [parsed row]} without catalogue matching."""
    table = read_table(data, filename)
    if not table:
        raise ValueError("The file is empty")
    hi = _header_index(table)
    headers = [_cell_text(c) or f"Column {j + 1}" for j, c in enumerate(table[hi])]
    title_lines = [" ".join(_cell_text(c) for c in r if _cell_text(c)) for r in table[:hi]]
    auto = detect_mapping(headers)
    m = {f: (mapping or {}).get(f, auto[f]) if mapping is not None else auto[f] for f in FIELDS}
    m = {f: (h if h in headers else None) for f, h in m.items()}
    idx = {f: headers.index(h) for f, h in m.items() if h}
    rows = []
    for n, r in enumerate(table[hi + 1:], start=hi + 2):
        if not any(_cell_text(c) for c in r):
            continue
        if len(rows) >= MAX_ROWS:
            raise ValueError(f"The list has more than {MAX_ROWS} rows")

        def cell(f):
            j = idx.get(f)
            return r[j] if j is not None and j < len(r) else None
        row = {"line": n}
        for f in ("product_id", "article_no", "manufacturer", "model", "description", "unit"):
            v = _cell_text(cell(f))
            row[f] = v or None
        lp, disc, net = _num(cell("list_price")), _num(cell("discount_pct")), _num(cell("net_price"))
        row["list_price"], row["discount_pct"] = lp, disc
        if net is None and lp is not None:
            net = round(lp * (1 - (disc or 0) / 100), 2)
            row["net_from"] = "list price - discount" if disc else "list price"
        row["net_price"] = net
        if net is None:
            row["problem"] = "No price in this row"
        elif net <= 0:
            row["problem"] = "Price is 0 or negative"
        elif not (row["product_id"] or row["model"]):
            row["problem"] = "No product id and no model to match"
        rows.append(row)
    return {"headers": headers, "mapping": m, "auto_mapping": auto, "title_lines": title_lines, "rows": rows}


# ---------------------------------------------------------------- matching to the catalogue

def match_rows(rows: list[dict], catalog: dict[str, dict], supplier: str | None) -> list[dict]:
    """Copies of the rows with matched_product_id, match_by, catalog_net and problem."""
    by_id = {pid.upper(): pid for pid in catalog}
    by_mm: dict[tuple[str, str], str] = {}
    by_model: dict[str, list[str]] = defaultdict(list)
    for pid, c in catalog.items():
        if c.get("model"):
            by_mm.setdefault((_fold(c.get("manufacturer")), _fold(c["model"])), pid)
            by_model[_fold(c["model"])].append(pid)
    seen: dict[str, int] = {}
    out = []
    for r in rows:
        r = dict(r)
        r.update(matched_product_id=None, match_by=None, catalog_net=None, problem=r.get("problem"))
        if r.get("problem"):
            out.append(r)
            continue
        pid, by = None, None
        if r.get("product_id"):
            pid, by = by_id.get(r["product_id"].strip().upper()), "product_id"
        if not pid and r.get("model"):
            mf = _fold(r.get("manufacturer") or supplier)
            pid, by = by_mm.get((mf, _fold(r["model"]))), "manufacturer+model"
            if not pid and len(by_model.get(_fold(r["model"]), [])) == 1:
                pid, by = by_model[_fold(r["model"])][0], "model"
        if not pid:
            r["problem"] = "Not in the catalogue"
            out.append(r)
            continue
        c = catalog[pid]
        lu, cu = normalize_unit(r.get("unit")), normalize_unit(c.get("unit")) or c.get("unit")
        if r.get("unit") and lu and cu and lu != cu:
            r.update(matched_product_id=pid, match_by=by, problem=f"Unit differs: list says {r['unit']}, catalogue {c.get('unit')}")
            out.append(r)
            continue
        if pid in seen:
            r.update(matched_product_id=pid, match_by=by, problem=f"Same product as row {seen[pid]}, which is used")
            out.append(r)
            continue
        seen[pid] = r["line"]
        r.update(matched_product_id=pid, match_by=by, catalog_net=float(c["net_cost_chf"]))
        out.append(r)
    return out


# ---------------------------------------------------------------- storage

def _dir() -> Path:
    return Path(config.PRICE_LISTS_DIR)


def _index_path() -> Path:
    return _dir() / "index.json"


_index_cache: tuple[tuple, list[dict]] | None = None


def version() -> tuple:
    """Changes whenever the lists, their order or their state change (part of the pricing cache key)."""
    try:
        return (str(_dir()), _index_path().stat().st_mtime_ns)
    except FileNotFoundError:
        return ()


def registry() -> list[dict]:
    """All lists, in priority order (first wins)."""
    global _index_cache
    v = version()
    if not v:
        return []
    if _index_cache and _index_cache[0] == v:
        return _index_cache[1]
    try:
        lists = json.loads(_index_path().read_text("utf-8")).get("lists") or []
    except (OSError, json.JSONDecodeError):
        lists = []
    _index_cache = (v, lists)
    return lists


def _save_registry(lists: list[dict]) -> None:
    global _index_cache
    _dir().mkdir(parents=True, exist_ok=True)
    tmp = _index_path().with_suffix(".tmp")
    tmp.write_text(json.dumps({"lists": lists}, ensure_ascii=False, indent=1), "utf-8")
    tmp.replace(_index_path())
    _index_cache = None
    _effective_cache.clear()


_rows_cache: dict[str, list[dict]] = {}


def _raw_rows(list_id: str) -> list[dict]:
    if list_id == "_preview":
        return _rows_cache.get(list_id, [])
    p = _dir() / f"{list_id}.json"
    if str(p) not in _rows_cache:
        _rows_cache[str(p)] = json.loads(p.read_text("utf-8"))["rows"] if p.exists() else []
    return _rows_cache[str(p)]


def rows_of(meta: dict, catalog: dict | None = None) -> list[dict]:
    return match_rows(_raw_rows(meta["id"]), catalog if catalog is not None else ref.load_catalog(),
                      meta.get("supplier"))


def get(list_id: str) -> dict:
    meta = next((m for m in registry() if m["id"] == list_id), None)
    if meta is None:
        raise _NoList(list_id)
    return meta


def state(meta: dict, today: date | None = None) -> str:
    """in_effect | inactive | not_yet_valid | expired"""
    today = today or _today()
    if not meta.get("active"):
        return "inactive"
    vf, vt = _date(meta.get("valid_from")), _date(meta.get("valid_to"))
    if vf and vf > today:
        return "not_yet_valid"
    if vt and vt < today:
        return "expired"
    return "in_effect"


def _rank(meta: dict) -> tuple:
    d = _date(meta.get("valid_from")) or date.min
    return (0 if meta.get("kind") == "negotiated" else 1, -d.toordinal(), meta.get("created_at") or "")


def default_order(lists: list[dict]) -> list[dict]:
    """Negotiated lists first, then supplier lists; within each, the most recent valid_from first."""
    return sorted(lists, key=_rank)


# ---------------------------------------------------------------- resolution

_variant: contextvars.ContextVar[dict | None] = contextvars.ContextVar("price_list_variant", default=None)
_effective_cache: dict[tuple, dict[str, dict]] = {}


@contextlib.contextmanager
def using(lists: list[dict]):
    """Price as if `lists` (in this order) were the registry, e.g. to show what activating one would do."""
    tok = _variant.set({"lists": lists, "catalog": None})
    try:
        yield
    finally:
        _variant.reset(tok)


def source_text(row: dict) -> str:
    """The human source of a list price: "Equans negotiated prices 2026 (valid from 1 Jan 2026), row 4"."""
    when = f" (valid from {fmt_date(row['price_list_date'])})" if row.get("price_list_date") else ""
    art = f", article {row['price_list_article_no']}" if row.get("price_list_article_no") else ""
    return f"{row['price_list']}{when}, row {row.get('price_list_row')}{art}"


def _build(base: dict[str, dict], lists: list[dict], today: date) -> dict[str, dict]:
    out = dict(base)
    taken: set[str] = set()
    for meta in lists:
        if state(meta, today) != "in_effect":
            continue
        for r in rows_of(meta, base):
            pid = r.get("matched_product_id")
            if r.get("problem") or not pid or pid in taken:
                continue
            taken.add(pid)
            c = base[pid]
            row = dict(c, net_cost_chf=float(r["net_price"]), catalog_net_cost_chf=float(c["net_cost_chf"]),
                       list_price_chf=r.get("list_price"), discount_pct=r.get("discount_pct"),
                       price_list=meta["name"], price_list_id=meta["id"], price_list_kind=meta.get("kind"),
                       price_list_date=meta.get("valid_from"), price_list_valid_to=meta.get("valid_to"),
                       supplier=meta.get("supplier"), price_list_row=r["line"],
                       price_list_article_no=r.get("article_no"))
            row["price_list_source"] = source_text(row)
            out[pid] = row
    return out


def effective_catalog() -> dict[str, dict]:
    """The catalogue with each product's net price resolved through the active price lists.

    Rows priced by a list carry price_list, price_list_id, price_list_date, supplier,
    price_list_source and catalog_net_cost_chf (the standard catalogue price)."""
    base = ref.load_catalog()
    var = _variant.get()
    if var is not None:
        if var["catalog"] is None:
            var["catalog"] = _build(base, var["lists"], _today())
        return var["catalog"]
    lists = registry()
    if not lists:
        return base
    key = (version(), id(base), _today())
    hit = _effective_cache.get(key)
    if hit is None:
        _effective_cache.clear()
        hit = _effective_cache[key] = _build(base, lists, _today())
    return hit


# ---------------------------------------------------------------- effect on tenders

def _tenders():
    from app import store
    for tid in store.list_ids():
        try:
            t = store.load(tid)
        except Exception:
            continue
        if t.get("status") in ("ready", "submitted", "awarded") and t.get("positions"):
            yield t


def _unit_prices(v: dict) -> dict[str, dict]:
    return {p["lv_id"]: p["pricing"] for p in v["positions"] if p.get("pricing")}


def compare(t: dict, before: list[dict], after: list[dict]) -> dict:
    """Direct cost, cost and the lines whose unit price moves, between two price list setups."""
    from app import pipeline
    with using(before):
        a = pipeline.view(t, use_cache=False)
    with using(after):
        b = pipeline.view(t, use_cache=False)
    pa, pb = _unit_prices(a), _unit_prices(b)
    lines = []
    for lv, x in pb.items():
        y = pa.get(lv)
        if y and abs(float(y["net_unit_cost"]) - float(x["net_unit_cost"])) >= 0.005:
            lines.append({"lv_id": lv, "product_id": x.get("product_id"), "old": y["net_unit_cost"],
                          "new": x["net_unit_cost"], "old_source": y["sources"].get("unit_price"),
                          "new_source": x["sources"].get("unit_price"),
                          "direct_delta": round(x["direct"] - y["direct"], 2)})
    ta, tb = a["totals"], b["totals"]
    return {"tender_id": t["id"], "name": t.get("name"), "status": t.get("status"),
            "direct_before": ta["direct"], "direct_after": tb["direct"],
            "direct_delta": round(tb["direct"] - ta["direct"], 2),
            "cost_delta": round(tb["cost"] - ta["cost"], 2),
            "lines_changed": len(lines), "lines": lines}


def effects(before: list[dict], after: list[dict], only_open: bool = True) -> list[dict]:
    out = []
    for t in _tenders():
        if only_open and t.get("status") != "ready":
            continue
        e = compare(t, before, after)
        e.pop("lines")
        out.append(e)
    return out


def _log_repricing(before: list[dict], after: list[dict], what: str, by: str) -> int:
    """One log entry per line whose unit price moved (old -> new, with sources), plus a summary."""
    from app import pipeline, store
    from app.pricing import swiss
    n = 0
    for t in list(_tenders()):
        e = compare(t, before, after)
        if not e["lines"]:
            continue
        batch = store.new_id("b_")
        with store.edit(t["id"], by=by) as tt:
            for ln in e["lines"]:
                le = store.append_log(tt, by, "price_list", target=("position", ln["lv_id"]),
                                      before={"unit_price": ln["old"], "unit_price_source": ln["old_source"]},
                                      after={"unit_price": ln["new"], "unit_price_source": ln["new_source"]},
                                      undoable=False, batch=batch,
                                      text=f"{what}: unit price CHF {swiss(ln['old'])} → CHF {swiss(ln['new'])}")
                le["direct_delta"] = ln["direct_delta"]
            sign = "+" if e["direct_delta"] >= 0 else "-"
            le = store.append_log(tt, by, "price_list", target=("tender", tt["id"]),
                             before={"direct": e["direct_before"]}, after={"direct": e["direct_after"]},
                             undoable=False, batch=batch,
                             text=(f"{what}: {e['lines_changed']} line{'s' if e['lines_changed'] != 1 else ''} re-priced, "
                                   f"direct cost {sign}CHF {swiss(abs(e['direct_delta']))}"))
            le["direct_delta"] = e["direct_delta"]
        pipeline._view_cache.pop(t["id"], None)
        n += 1
    return n


# ---------------------------------------------------------------- import, preview, changes

def _suggest(filename: str, title_lines: list[str], rows: list[dict]) -> dict:
    stem = Path(filename or "price list").stem
    words = re.sub(r"[_\-]+", " ", stem).strip()
    low = words.lower()
    kind = "negotiated" if re.search(r"equans|negotiat|konditio|rahmen|eigene", low) else "supplier"
    title = next((x for x in title_lines if x.strip()), "")
    name = re.split(r"\s+[·|–-]\s+gültig|\s+[·|]\s+valid", title, flags=re.I)[0].strip() if title else ""
    name = name or (words[:1].upper() + words[1:])
    supplier = None
    if title:
        m = re.match(r"^(.+?\b(?:AG|GmbH|SA|Sàrl|Ltd|Inc)\b)", title)
        supplier = m[1].strip() if m else None
    if not supplier:
        rest = [w for w in words.split() if not re.fullmatch(
            r"(?i)supplier|lieferant|pricelist|preisliste|price|prices|preise|list|liste|net|netto|negotiated|\d{4}", w)]
        supplier = " ".join(w[:1].upper() + w[1:] for w in rest) or None
    if not supplier:
        mf = Counter(r.get("manufacturer") for r in rows if r.get("manufacturer"))
        supplier = mf.most_common(1)[0][0] if mf else None
    vf = None
    for x in title_lines + [words]:
        m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", x) or re.search(r"(\d{4})-(\d{2})-(\d{2})", x)
        if m:
            vf = (f"{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}" if "." in m[0] else m[0])
            break
    if not vf:
        y = re.search(r"\b(20\d{2})\b", words)
        vf = f"{y[1]}-01-01" if y else _today().isoformat()
    return {"name": name, "supplier": supplier, "kind": kind, "valid_from": vf, "valid_to": None, "currency": "CHF"}


def _clean_meta(meta: dict, fallback: dict) -> dict:
    m = {**fallback, **{k: v for k, v in (meta or {}).items() if v not in (None, "")}}
    if m.get("kind") not in KINDS:
        raise ValueError("Kind must be 'negotiated' or 'supplier'")
    if str(m.get("currency") or "CHF").upper() != "CHF":
        raise ValueError("Only CHF price lists are supported")
    vf, vt = _date(m.get("valid_from")), _date(m.get("valid_to"))
    if vf and vt and vt < vf:
        raise ValueError("Valid to is before valid from")
    if not str(m.get("name") or "").strip():
        raise ValueError("Give the price list a name")
    return {"name": str(m["name"]).strip(), "supplier": (str(m.get("supplier") or "").strip() or None),
            "kind": m["kind"], "valid_from": vf.isoformat() if vf else None,
            "valid_to": vt.isoformat() if vt else None, "currency": "CHF"}


def _counts(rows: list[dict]) -> dict:
    ok = [r for r in rows if not r.get("problem")]
    return {"rows": len(rows), "matched": len(ok), "unmatched": len(rows) - len(ok),
            "cheaper": sum(1 for r in ok if r["net_price"] < r["catalog_net"] - 0.005),
            "dearer": sum(1 for r in ok if r["net_price"] > r["catalog_net"] + 0.005)}


def _insert(lists: list[dict], meta: dict) -> list[dict]:
    """Insert at the default position (before the first list it outranks), keeping a custom order."""
    i = next((k for k, m in enumerate(lists) if _rank(meta) < _rank(m)), len(lists))
    return lists[:i] + [meta] + lists[i:]


def preview(data: bytes, filename: str, mapping: dict | None = None, meta: dict | None = None) -> dict:
    parsed = parse_table(data, filename, mapping)
    sugg = _suggest(filename, parsed["title_lines"], parsed["rows"])
    m = {**sugg, **{k: v for k, v in (meta or {}).items() if v not in (None, "")}}
    rows = match_rows(parsed["rows"], ref.load_catalog(), m.get("supplier"))
    has_price = bool(parsed["mapping"].get("net_price") or parsed["mapping"].get("list_price"))
    has_key = bool(parsed["mapping"].get("product_id") or parsed["mapping"].get("model"))
    missing = ([] if has_price else ["net_price"]) + ([] if has_key else ["product_id"])
    effect = []
    if not missing and m.get("kind") in KINDS:
        try:
            draft = {**_clean_meta(m, sugg), "id": "_preview", "active": True}
            _rows_cache["_preview"] = parsed["rows"]
            cur = registry()
            effect = effects(cur, _insert(cur, draft))
        except ValueError:
            effect = []
        finally:
            _rows_cache.pop("_preview", None)
    return {"file_name": filename, "headers": parsed["headers"], "mapping": parsed["mapping"],
            "auto_mapping": parsed["auto_mapping"], "title_lines": parsed["title_lines"],
            "fields": [{"key": f, "label": FIELD_LABELS[f]} for f in FIELDS], "missing": missing,
            "suggested": sugg, "meta": m, "rows": rows, "counts": _counts(rows), "effect": effect}


def create(data: bytes, filename: str, mapping: dict | None, meta: dict | None, activate: bool = True,
           by: str = "You") -> dict:
    parsed = parse_table(data, filename, mapping)
    if not (parsed["mapping"].get("net_price") or parsed["mapping"].get("list_price")):
        raise ValueError("Pick the column with the net price (or the list price)")
    if not (parsed["mapping"].get("product_id") or parsed["mapping"].get("model")):
        raise ValueError("Pick the column with our product id or the model / type")
    sugg = _suggest(filename, parsed["title_lines"], parsed["rows"])
    clean = _clean_meta(meta or {}, sugg)
    lid = "pl_" + secrets.token_hex(3)
    ext = Path(filename or "").suffix.lower() or ".csv"
    _dir().mkdir(parents=True, exist_ok=True)
    (_dir() / f"{lid}{ext}").write_bytes(data)
    (_dir() / f"{lid}.json").write_text(json.dumps({"mapping": parsed["mapping"], "rows": parsed["rows"]},
                                                   ensure_ascii=False), "utf-8")
    rows = match_rows(parsed["rows"], ref.load_catalog(), clean["supplier"])
    entry = {"id": lid, **clean, "file": f"{lid}{ext}", "original_name": filename, "active": False,
             "created_at": datetime.now().isoformat(timespec="seconds"), "created_by": by,
             "rows": len(rows), "matched": _counts(rows)["matched"], "mapping": parsed["mapping"]}
    before = registry()
    _save_registry(_insert(before, entry))
    if activate:
        update(lid, {"active": True}, by=by)
    return get(lid)


def update(list_id: str, patch: dict, by: str = "You") -> dict:
    """Activate / deactivate, rename or change the dates of a list. Re-prices and logs changed lines."""
    before = registry()
    meta = get(list_id)
    new = dict(meta)
    if "active" in patch and patch["active"] is not None:
        new["active"] = bool(patch["active"])
    for k in ("name", "supplier", "kind", "valid_from", "valid_to"):
        if k in patch and patch[k] is not None:
            new[k] = patch[k] or None
    new.update(_clean_meta({k: new.get(k) for k in ("name", "supplier", "kind", "valid_from", "valid_to")}, {}))
    if new == meta:
        return meta
    after = [new if m["id"] == list_id else m for m in before]
    if new["active"] != meta["active"]:
        what = f"{new['name']} {'activated' if new['active'] else 'deactivated'}"
    else:
        what = f"{new['name']} changed"
    _save_registry(after)
    _log_repricing(before, after, what, by)
    return new


def reorder(order: list[str], by: str = "You") -> list[dict]:
    """New priority order (list ids, first wins). "default" sorts by the default rule."""
    before = registry()
    if order == ["default"]:
        after = default_order(before)
    else:
        ids = [m["id"] for m in before]
        if sorted(order) != sorted(ids):
            raise ValueError("The order must name every price list once")
        after = [next(m for m in before if m["id"] == i) for i in order]
    if [m["id"] for m in after] == [m["id"] for m in before]:
        return before
    _save_registry(after)
    _log_repricing(before, after, "Price list order changed", by)
    return after


def delete(list_id: str, by: str = "You") -> None:
    before = registry()
    meta = get(list_id)
    after = [m for m in before if m["id"] != list_id]
    _save_registry(after)
    _log_repricing(before, after, f"{meta['name']} deleted", by)
    for p in _dir().glob(f"{list_id}.*"):
        _rows_cache.pop(str(p), None)
        p.unlink(missing_ok=True)


def summary(meta: dict, with_effect: bool = True) -> dict:
    """A list with its counts, its state today, the products it prices and the effect of switching it."""
    lists = registry()
    rows = rows_of(meta)
    eff = effective_catalog()
    winning = [r for r in rows if not r.get("problem") and eff.get(r["matched_product_id"], {}).get("price_list_id") == meta["id"]]
    out = {**meta, "state": state(meta), "position": next((i for i, m in enumerate(lists) if m["id"] == meta["id"]), None),
           "counts": {**_counts(rows), "in_use": len(winning),
                      "changes": sum(1 for r in winning if abs(r["net_price"] - r["catalog_net"]) >= 0.005)}}
    if with_effect:
        # what this list does to the open tenders: compared with the same setup without it (active) or with it (inactive)
        without = [m for m in lists if m["id"] != meta["id"]]
        withit = [dict(m, active=True) if m["id"] == meta["id"] else m for m in lists]
        out["effect"] = effects(without, withit)
        out["effect_kind"] = "saves" if meta.get("active") else "would"
    return out


def detail(list_id: str) -> dict:
    meta = get(list_id)
    out = summary(meta)
    eff = effective_catalog()
    cat = ref.load_catalog()
    rows = []
    for r in rows_of(meta):
        pid = r.get("matched_product_id")
        c = cat.get(pid) if pid else None
        used = eff.get(pid) if pid else None
        rows.append({**r, "catalog_description": (c or {}).get("description"), "catalog_unit": (c or {}).get("unit"),
                     "catalog_model": " ".join(x for x in ((c or {}).get("manufacturer"), (c or {}).get("model")) if x) or None,
                     "in_use": bool(used and used.get("price_list_id") == list_id and not r.get("problem")),
                     "used_price": (used or {}).get("net_cost_chf") if used else None,
                     "used_source": ((used or {}).get("price_list") or "Catalogue") if used else None})
    out["rows_detail"] = rows
    return out


def overview() -> dict:
    lists = registry()
    return {"lists": [summary(m) for m in lists], "precedence": PRECEDENCE,
            "kinds": KIND_LABELS, "demo_files": demo_files()}


def demo_files() -> list[dict]:
    d = config.DATA_DIR / "demo"
    return [{"name": p.name} for p in sorted(d.glob("*")) if p.suffix.lower() in (".csv", ".xlsx")
            and re.search(r"price", p.name, re.I)]


def demo_file(name: str) -> tuple[bytes, str]:
    files = {f["name"] for f in demo_files()}
    if name not in files:
        raise ValueError(f"No demo price list named {name}")
    return (config.DATA_DIR / "demo" / name).read_bytes(), name
