"""Read a table out of whatever file a supplier or Equans sends: one reader for the catalogue,
the install rules and the price lists.

Formats (detected by magic bytes first, then by the file name):
  text   CSV, TSV, semicolon or pipe separated (UTF-8, UTF-16 with BOM, Windows cp1252)
  .xlsx / .xlsm (openpyxl), legacy .xls (xlrd), .ods (pandas + odfpy)
  .json  an array of objects (or of arrays), or {"rows" | "catalog" | ...: [...]}
  .pdf   tables found by PyMuPDF on every page; a table whose header repeats on the next page
         is joined to the one before

``read_table`` returns the rows as lists of cell values (str, numbers, dates or None). When a
``score`` function is given, the sheet (or PDF table) whose rows score highest is used, so a
workbook with a cover sheet still finds the price sheet.
"""
from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import zipfile
from datetime import date, datetime
from typing import Callable

MAX_ROWS = 20000
FORMATS_TEXT = "Excel (.xlsx, .xls), OpenDocument (.ods), CSV/TSV, JSON or a PDF with a table"

Score = Callable[[list[list]], float]


# ---------------------------------------------------------------- cells

def fold(s) -> str:
    """Header comparison key: lower case, umlauts and accents folded, no spaces or punctuation."""
    s = str(s or "").lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        s = s.replace(a, b)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]", "", s)


def num(v, pct_cell: bool = False) -> float | None:
    """Numbers as suppliers write them: 1'234.50, 1.234,50, 74,80, CHF 74.80, 12 %."""
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return round(float(v) * (100 if pct_cell else 1), 6)
    s = re.sub(r"(?i)chf|fr\.|sfr\.?|%|\s", "", str(v)).replace("’", "'").replace("'", "").replace("`", "")
    if not s or s in ("-", "–"):
        return None
    if s.endswith((".-", ".–")):  # "12.-" Swiss whole francs
        s = s[:-2]
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif s.count(",") == 1:
        s = s.replace(",", ".")
    elif s.count(",") > 1:
        s = s.replace(",", "")
    elif s.count(".") > 1:  # 1.234.567 thousands dots
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v != v:  # NaN from pandas
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


# ---------------------------------------------------------------- format detection

def detect_format(data: bytes, filename: str | None) -> str:
    """xlsx | ods | xls | pdf | json | text, from the content first and the extension second."""
    name = (filename or "").lower()
    head = data[:8]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if head.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = set(z.namelist())
                if "mimetype" in names and b"opendocument.spreadsheet" in z.read("mimetype"):
                    return "ods"
                if "content.xml" in names and not any(n.startswith("xl/") for n in names):
                    return "ods"
        except zipfile.BadZipFile:
            pass
        return "xlsx"
    if name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if name.endswith(".xls"):
        return "xls"
    if name.endswith(".ods"):
        return "ods"
    if name.endswith(".pdf"):
        return "pdf"
    if name.endswith(".json") or data.lstrip(b"\xef\xbb\xbf \t\r\n")[:1] in (b"[", b"{"):
        return "json"
    if name.endswith((".numbers", ".xlsb", ".doc", ".docx")):
        raise ValueError(f"This file type is not supported. Use {FORMATS_TEXT}")
    return "text"


# ---------------------------------------------------------------- readers

def _pick(candidates: list[list[list]], score: Score | None) -> list[list]:
    candidates = [c for c in candidates if any(any(cell_text(x) for x in r) for r in c)]
    if not candidates:
        return []
    if score is None:
        return candidates[0]
    best = max(range(len(candidates)), key=lambda i: (score(candidates[i]), -i))
    return candidates[best]


def _xlsx(data: bytes) -> list[list[list]]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # not a readable workbook
        raise ValueError(f"Could not read the Excel file: {e}") from None
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows():
            # percentage cells come back as percent numbers (0.12 formatted 12% -> 12.0)
            rows.append([(round(c.value * 100, 6) if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
                          and "%" in str(getattr(c, "number_format", "") or "") else c.value) for c in r])
            if len(rows) > MAX_ROWS + 20:
                break
        sheets.append(rows)
    wb.close()
    return sheets


def _xls(data: bytes) -> list[list[list]]:
    try:
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
    except Exception as e:
        raise ValueError(f"Could not read the old Excel (.xls) file: {e}") from None
    sheets = []
    for sh in book.sheets():
        rows = []
        for i in range(min(sh.nrows, MAX_ROWS + 21)):
            out = []
            for c in sh.row(i):
                if c.ctype == xlrd.XL_CELL_DATE:
                    try:
                        out.append(xlrd.xldate.xldate_as_datetime(c.value, book.datemode))
                    except Exception:
                        out.append(c.value)
                elif c.ctype == xlrd.XL_CELL_BOOLEAN:
                    out.append(bool(c.value))
                elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    out.append(None)
                else:
                    out.append(c.value)
            rows.append(out)
        sheets.append(rows)
    return sheets


def _ods(data: bytes) -> list[list[list]]:
    try:
        import pandas as pd
        frames = pd.read_excel(io.BytesIO(data), engine="odf", sheet_name=None, header=None, dtype=object)
    except Exception as e:
        raise ValueError(f"Could not read the OpenDocument (.ods) file: {e}") from None
    sheets = []
    for df in frames.values():
        rows = []
        for r in df.itertuples(index=False):
            rows.append([None if (v is None or (isinstance(v, float) and v != v)) else
                         (v.to_pydatetime() if hasattr(v, "to_pydatetime") else v) for v in r])
        sheets.append(rows[:MAX_ROWS + 21])
    return sheets


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    if len(data) > 3 and data[1:2] == b"\x00" and data[3:4] == b"\x00":
        try:
            return data.decode("utf-16-le")
        except UnicodeDecodeError:
            pass
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Could not read the file as text. Use {FORMATS_TEXT}")


def _text(data: bytes) -> list[list[list]]:
    text = _decode(data)
    if "\x00" in text[:2000]:
        raise ValueError(f"This file type is not supported. Use {FORMATS_TEXT}")
    lines = text.splitlines()[:20]
    sample = "\n".join(lines)
    counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
    # a delimiter that appears on every non-empty line beats one that only shows up in numbers ("74,80")
    steady = {d: n for d, n in counts.items() if n and all(d in ln for ln in lines if ln.strip())}
    pool = steady or counts
    delim = max(pool, key=pool.get) if any(pool.values()) else ","
    return [[r for r in csv.reader(io.StringIO(text), delimiter=delim)]]


ROW_KEYS = ("rows", "catalog", "catalogue", "products", "items", "data", "rules", "lines", "prices",
            "price_list", "pricelist", "artikel", "produkte")


def _json(data: bytes) -> list[list[list]]:
    try:
        obj = json.loads(_decode(data))
    except json.JSONDecodeError as e:
        raise ValueError(f"The file is not valid JSON: {e}") from None
    items = obj
    if isinstance(obj, dict):
        items = next((obj[k] for k in ROW_KEYS if isinstance(obj.get(k), list)), None)
        if items is None:
            lists = [v for v in obj.values() if isinstance(v, list)]
            items = lists[0] if len(lists) == 1 else None
    if not isinstance(items, list):
        raise ValueError('JSON must be an array of rows, or an object with a "rows" array')
    if all(isinstance(x, list) for x in items):
        return [items]
    if not all(isinstance(x, dict) for x in items):
        raise ValueError("JSON rows must all be objects (or all arrays)")
    headers: list[str] = []
    for x in items:
        for k in x:
            if k not in headers:
                headers.append(k)
    rows: list[list] = [list(headers)]
    for x in items:
        rows.append([("|".join(map(str, v)) if isinstance(v, list) else v) for v in (x.get(h) for h in headers)])
    return [rows]


def _pdf_rows(page, table) -> list[list]:
    """Cell texts read with the page's own text extraction (find_tables' extract() splits underscores
    and other low glyphs onto a line of their own: FIRE_DAMPER -> "FIRE DAMPER _")."""
    import pymupdf
    out = []
    for row in table.rows:
        cells = []
        for bbox in row.cells:
            if bbox is None:
                cells.append(None)
                continue
            t = page.get_text("text", clip=pymupdf.Rect(bbox))
            cells.append(re.sub(r"\s*\n\s*", " ", t).strip())
        out.append(cells)
    return [r for r in out if any(cell_text(c) for c in r)]


def _pdf(data: bytes, what: str = "list") -> list[list[list]]:
    try:
        import pymupdf
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:
        raise ValueError(f"Could not read the PDF: {e}") from None
    groups: list[list[list]] = []   # each: header row + data rows, tables joined across pages
    try:
        for page in doc:
            found = page.find_tables()
            for tb in found.tables:
                rows = _pdf_rows(page, tb)
                if not rows:
                    continue
                key = [fold(c) for c in rows[0]]
                same = next((g for g in groups if [fold(c) for c in g[0]] == key), None)
                if same is not None:            # the header repeats on the next page
                    same.extend(rows[1:])
                elif groups and len(rows[0]) == len(groups[-1][0]) and any(num(c) is not None for c in rows[0]):
                    groups[-1].extend(rows)     # same columns, first row is data: the table continues
                else:
                    groups.append(rows)
                if sum(len(g) for g in groups) > MAX_ROWS + 20:
                    break
    finally:
        doc.close()
    if not groups:
        raise ValueError(f"No table found in this PDF. Export the {what} as Excel or CSV")
    return groups


def read_table(data: bytes, filename: str | None, score: Score | None = None, what: str = "list") -> list[list]:
    """All rows of the best sheet / table of the file, as lists of cell values."""
    rows, _fmt = read_table_fmt(data, filename, score, what)
    return rows


def read_table_fmt(data: bytes, filename: str | None, score: Score | None = None,
                   what: str = "list") -> tuple[list[list], str]:
    """``read_table`` plus the detected format (xlsx, xls, ods, pdf, json or csv)."""
    if not data:
        raise ValueError("The file is empty")
    fmt = detect_format(data, filename)
    if fmt == "pdf":
        sheets = _pdf(data, what)
    else:
        sheets = {"xlsx": _xlsx, "xls": _xls, "ods": _ods, "json": _json, "text": _text}[fmt](data)
    rows = _pick(sheets, score)
    return rows, ("csv" if fmt == "text" else fmt)
